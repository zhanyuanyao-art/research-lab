"""Dataset import, profiling, preview, rename/cast and delete.

Readers: CSV/TSV (encoding sniffed: UTF-8, UTF-8-BOM, GB18030 ⊇ GBK, Big5, Latin-1), XLSX (sheet
select), Stata .dta and SPSS .sav via pyreadstat, Parquet via DuckDB. Column names are kept as
written, Chinese included; blanks and duplicates get a suffix so every column is addressable."""
from __future__ import annotations

import csv
import io
import json
import shutil
import sqlite3
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd

from ..db.connection import next_sequence_id, now_iso
from . import store

SOURCE_KINDS = ("csv", "xlsx", "dta", "sav", "parquet")
EXT_KIND = {".csv": "csv", ".tsv": "csv", ".txt": "csv", ".xlsx": "xlsx", ".xlsm": "xlsx", ".dta": "dta",
            ".sav": "sav", ".zsav": "sav", ".parquet": "parquet", ".pq": "parquet"}
ENCODINGS = ("utf-8-sig", "gb18030", "big5", "latin-1")
CAST_TYPES = {"number": ("数值", "DOUBLE"), "integer": ("整数", "BIGINT"), "text": ("文本", "VARCHAR"),
              "date": ("日期", "DATE"), "bool": ("布尔", "BOOLEAN")}
PREVIEW_ROWS = 50
TOP_VALUES = 5
MAX_BYTES = 500 * 1024 * 1024


class DatasetError(ValueError):
    pass


def kind_of(filename: str) -> str:
    kind = EXT_KIND.get(Path(filename).suffix.lower())
    if not kind:
        raise DatasetError(f"不支持的文件类型：{Path(filename).suffix or '无扩展名'}（支持 CSV/TSV、XLSX、DTA、SAV、Parquet）")
    return kind


# ---------------------------------------------------------------- readers

def sniff_encoding(raw: bytes) -> str:
    sample = raw[: 1 << 20]
    for enc in ENCODINGS[:-1]:
        try:
            sample.decode(enc)
            return enc
        except UnicodeDecodeError as e:
            # a multi-byte character cut at the sample boundary is not a wrong guess
            if len(sample) == (1 << 20) and e.start >= len(sample) - 4:
                return enc
            continue
    return "latin-1"


def sniff_delimiter(text: str, filename: str) -> str:
    if Path(filename).suffix.lower() == ".tsv":
        return "\t"
    head = "\n".join(text.splitlines()[:20])
    try:
        return csv.Sniffer().sniff(head, delimiters=",\t;|").delimiter
    except csv.Error:
        return ","


def clean_columns(cols: list[Any]) -> list[str]:
    out: list[str] = []
    seen: dict[str, int] = {}
    for i, c in enumerate(cols):
        name = "" if c is None or (isinstance(c, float) and np.isnan(c)) else str(c).strip()
        if not name or name.lower().startswith("unnamed:"):
            name = f"col_{i + 1}"
        base = name
        if name in seen:
            seen[base] += 1
            name = f"{base}_{seen[base]}"
            while name in seen:
                seen[base] += 1
                name = f"{base}_{seen[base]}"
        seen.setdefault(name, 1)
        out.append(name)
    return out


def xlsx_sheets(path: Path) -> list[str]:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def read_file(path: Path, kind: str, *, original_name: str | None = None, sheet: str | None = None) -> tuple[pd.DataFrame, dict[str, Any]]:
    """Return the frame and reader metadata (encoding, delimiter, sheet, sheets, variable labels)."""
    meta: dict[str, Any] = {}
    name = original_name or path.name
    if kind == "csv":
        raw = path.read_bytes()
        enc = sniff_encoding(raw)
        text = raw.decode(enc, errors="replace")
        delim = sniff_delimiter(text, name)
        df = pd.read_csv(io.StringIO(text), sep=delim, low_memory=False)
        meta.update(encoding=enc, delimiter=delim)
    elif kind == "xlsx":
        sheets = xlsx_sheets(path)
        if not sheets:
            raise DatasetError("工作簿里没有工作表")
        use = sheet if sheet in sheets else sheets[0]
        df = pd.read_excel(path, sheet_name=use, engine="openpyxl")
        meta.update(sheet=use, sheets=sheets)
    elif kind in ("dta", "sav"):
        import pyreadstat
        reader = pyreadstat.read_dta if kind == "dta" else pyreadstat.read_sav
        df, rmeta = reader(str(path), apply_value_formats=False)
        labels = dict(zip(rmeta.column_names, rmeta.column_labels or []))
        meta["labels"] = {k: v for k, v in labels.items() if v}
        vl = {}
        for col, lblset in (rmeta.variable_to_label or {}).items():
            vals = (rmeta.value_labels or {}).get(lblset) or {}
            vl[col] = {str(k): str(v) for k, v in list(vals.items())[:50]}
        if vl:
            meta["value_labels"] = vl
    elif kind == "parquet":
        df = store.read_parquet_file(path)
    else:
        raise DatasetError(f"unknown source kind {kind}")
    df.columns = clean_columns(list(df.columns))
    if "labels" in meta:
        # labels were keyed by the raw names; cleaned names only differ for blanks/duplicates
        meta["labels"] = {k: v for k, v in meta["labels"].items() if k in df.columns}
    # object columns that are really numeric (e.g. from xlsx) become numbers
    for c in df.columns:
        if df[c].dtype == object:
            conv = pd.to_numeric(df[c], errors="coerce")
            if conv.notna().sum() == df[c].notna().sum() and df[c].notna().any():
                df[c] = conv
    return df, meta


# ---------------------------------------------------------------- profile

def _num(v: Any) -> float | int | None:
    if v is None:
        return None
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(v, (np.integer,)):
        return int(v)
    if isinstance(v, (np.floating, float)):
        return round(float(v), 6)
    if isinstance(v, (int,)):
        return v
    return None


def dtype_label(s: pd.Series) -> str:
    if pd.api.types.is_bool_dtype(s):
        return "bool"
    if pd.api.types.is_integer_dtype(s):
        return "integer"
    if pd.api.types.is_float_dtype(s):
        # float columns holding only whole numbers (ints with NaN) read as integer
        nn = s.dropna()
        return "integer" if len(nn) and bool((nn == nn.round()).all()) else "number"
    if pd.api.types.is_datetime64_any_dtype(s):
        return "date"
    return "text"


def profile(df: pd.DataFrame, labels: dict[str, str] | None = None) -> dict[str, Any]:
    n = int(len(df))
    cols = []
    for c in df.columns:
        s = df[c]
        missing = int(s.isna().sum())
        kind = dtype_label(s)
        entry: dict[str, Any] = {"name": str(c), "dtype": kind, "missing": missing,
                                 "missing_pct": round(100.0 * missing / n, 1) if n else 0.0,
                                 "unique": int(s.nunique(dropna=True))}
        if labels and labels.get(c):
            entry["label"] = labels[c]
        if kind in ("integer", "number") and not pd.api.types.is_bool_dtype(s):
            nn = s.dropna()
            if len(nn):
                entry.update(min=_num(nn.min()), max=_num(nn.max()), mean=_num(nn.mean()),
                             std=_num(nn.std()) if len(nn) > 1 else None, median=_num(nn.median()))
        if kind not in ("number",) or entry["unique"] <= 12:
            vc = s.dropna().astype(str).value_counts().head(TOP_VALUES)
            entry["top"] = [[str(k), int(v)] for k, v in vc.items()]
        cols.append(entry)
    return {"n": n, "k": int(len(df.columns)), "columns": cols,
            "complete_rows": int(df.dropna().shape[0]) if len(df.columns) else n}


def schema_of(df: pd.DataFrame, meta: dict[str, Any]) -> dict[str, Any]:
    labels = meta.get("labels") or {}
    out = {"columns": [{"name": str(c), "dtype": dtype_label(df[c]), **({"label": labels[c]} if labels.get(c) else {})}
                       for c in df.columns]}
    for k in ("encoding", "delimiter", "sheet", "sheets", "value_labels"):
        if k in meta:
            out[k] = meta[k]
    return out


# ---------------------------------------------------------------- registry

def _row(conn: sqlite3.Connection, ds_id: str) -> sqlite3.Row | None:
    return conn.execute("SELECT * FROM datasets WHERE id = ?", (ds_id,)).fetchone()


def get(conn: sqlite3.Connection, ds_id: str) -> dict[str, Any] | None:
    r = _row(conn, ds_id)
    if r is None:
        return None
    d = dict(r)
    d["schema"] = json.loads(r["schema_json"] or "{}")
    d["profile"] = json.loads(r["profile_json"] or "{}")
    d["columns"] = [c["name"] for c in d["schema"].get("columns", [])]
    return d


def list_all(conn: sqlite3.Connection, paper_id: str | None = None) -> list[dict[str, Any]]:
    sql, args = ("SELECT d.*, p.title AS paper_title, (SELECT COUNT(*) FROM analysis_runs r WHERE r.dataset_id = d.id) AS n_runs "
                 "FROM datasets d LEFT JOIN papers p ON p.id = d.paper_id"), []
    if paper_id:
        sql += " WHERE d.paper_id = ?"; args.append(paper_id)
    sql += " ORDER BY d.created_at DESC, d.id DESC"
    out = []
    for r in conn.execute(sql, args):
        d = dict(r)
        d["columns"] = [c["name"] for c in json.loads(r["schema_json"] or "{}").get("columns", [])]
        out.append(d)
    return out


def column_kinds(ds: dict[str, Any]) -> dict[str, str]:
    return {c["name"]: c["dtype"] for c in ds["schema"].get("columns", [])}


def _refresh(conn: sqlite3.Connection, ds_id: str, df: pd.DataFrame, meta: dict[str, Any]) -> None:
    conn.execute("UPDATE datasets SET rows=?, cols=?, schema_json=?, profile_json=?, updated_at=? WHERE id=?",
                 (int(len(df)), int(len(df.columns)), json.dumps(schema_of(df, meta), ensure_ascii=False),
                  json.dumps(profile(df, meta.get("labels")), ensure_ascii=False), now_iso(), ds_id))


def import_file(conn: sqlite3.Connection, src: Path | str, *, original_name: str | None = None, name: str | None = None,
                paper_id: str | None = None, direction: str | None = None, note: str | None = None,
                sheet: str | None = None) -> str:
    """Copy the original, load it into DuckDB + parquet, register the dataset. Returns DS-YYYY-NNNN."""
    src = Path(src)
    original = original_name or src.name
    kind = kind_of(original)
    if src.stat().st_size > MAX_BYTES:
        raise DatasetError("文件超过 500 MB，请先在本地拆分或转为 Parquet")
    if paper_id and not conn.execute("SELECT 1 FROM papers WHERE id = ?", (paper_id,)).fetchone():
        raise DatasetError(f"没有这篇论文：{paper_id}")
    try:
        df, meta = read_file(src, kind, original_name=original, sheet=sheet)
    except DatasetError:
        raise
    except Exception as e:  # noqa: BLE001 - reader errors are shown to the user
        raise DatasetError(f"无法读取文件：{type(e).__name__}: {e}") from e
    if df.shape[1] == 0:
        raise DatasetError("文件里没有可读取的列")
    ds_id = next_sequence_id(conn, "datasets", "DS")
    folder = store.dataset_dir(ds_id)
    folder.mkdir(parents=True, exist_ok=True)
    safe = Path(original).name.replace("/", "_") or f"original.{kind}"
    shutil.copy2(src, folder / safe)
    store.write_df(ds_id, df)
    now = now_iso()
    conn.execute(
        "INSERT INTO datasets(id, name, source_kind, original_name, rows, cols, schema_json, profile_json, paper_id, direction, note, created_at, updated_at) "
        "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (ds_id, (name or Path(original).stem).strip() or ds_id, kind, safe, int(len(df)), int(len(df.columns)),
         json.dumps(schema_of(df, meta), ensure_ascii=False), json.dumps(profile(df, meta.get("labels")), ensure_ascii=False),
         paper_id or None, direction or None, note or None, now, now))
    return ds_id


def original_path(ds: dict[str, Any]) -> Path:
    return store.dataset_dir(ds["id"]) / ds["original_name"]


def reimport_sheet(conn: sqlite3.Connection, ds_id: str, sheet: str) -> None:
    ds = get(conn, ds_id)
    if not ds:
        raise DatasetError(f"没有这个数据集：{ds_id}")
    if ds["source_kind"] != "xlsx":
        raise DatasetError("只有 XLSX 数据集可以切换工作表")
    df, meta = read_file(original_path(ds), "xlsx", original_name=ds["original_name"], sheet=sheet)
    store.write_df(ds_id, df)
    _refresh(conn, ds_id, df, meta)


def preview(ds_id: str, limit: int = PREVIEW_ROWS) -> tuple[list[str], list[list[Any]]]:
    df = store.read_df(ds_id, limit=limit)
    rows = []
    for rec in df.itertuples(index=False):
        rows.append([None if (not isinstance(v, str) and pd.isna(v)) else (_num(v) if isinstance(v, (np.floating, float, np.integer)) else v)
                     for v in rec])
    return [str(c) for c in df.columns], rows


def update_meta(conn: sqlite3.Connection, ds_id: str, *, name: str | None = None, note: str | None = None,
                paper_id: str | None = None, direction: str | None = None) -> None:
    if not _row(conn, ds_id):
        raise DatasetError(f"没有这个数据集：{ds_id}")
    if paper_id and not conn.execute("SELECT 1 FROM papers WHERE id = ?", (paper_id,)).fetchone():
        raise DatasetError(f"没有这篇论文：{paper_id}")
    conn.execute("UPDATE datasets SET name = COALESCE(NULLIF(?, ''), name), note = ?, paper_id = ?, direction = ?, updated_at = ? WHERE id = ?",
                 ((name or "").strip(), (note or "").strip() or None, paper_id or None, direction or None, now_iso(), ds_id))


def rename_cast(conn: sqlite3.Connection, ds_id: str, renames: dict[str, str], casts: dict[str, str]) -> dict[str, int]:
    """Rename columns and cast types in one DuckDB statement. Failed casts become NULL (TRY_CAST);
    the returned dict counts how many previously non-missing values each cast turned into NULL."""
    ds = get(conn, ds_id)
    if not ds:
        raise DatasetError(f"没有这个数据集：{ds_id}")
    cols = ds["columns"]
    renames = {k: v.strip() for k, v in renames.items() if k in cols and v and v.strip() and v.strip() != k}
    casts = {k: v for k, v in casts.items() if k in cols and v in CAST_TYPES}
    final = [renames.get(c, c) for c in cols]
    if len(set(final)) != len(final):
        raise DatasetError("重命名后出现重复列名")
    if not renames and not casts:
        return {}
    table = store.table_name(ds_id)
    lost: dict[str, int] = {}
    with store.connect(read_only=True) as con:
        for c, t in casts.items():
            sql_t = CAST_TYPES[t][1]
            row = con.execute(f"SELECT COUNT(*) FILTER (WHERE {store.qi(c)} IS NOT NULL AND TRY_CAST({store.qi(c)} AS {sql_t}) IS NULL) "
                              f"FROM {table}").fetchone()
            lost[renames.get(c, c)] = int(row[0])
    parts = []
    for c in cols:
        expr = store.qi(c)
        if c in casts:
            expr = f"TRY_CAST({expr} AS {CAST_TYPES[casts[c]][1]})"
        parts.append(f"{expr} AS {store.qi(renames.get(c, c))}")
    store.execute_replace(ds_id, f"SELECT {', '.join(parts)} FROM {table}")
    df = store.read_df(ds_id)
    old_schema = ds["schema"]
    labels = {renames.get(c["name"], c["name"]): c["label"] for c in old_schema.get("columns", []) if c.get("label")}
    meta = {k: old_schema[k] for k in ("encoding", "delimiter", "sheet", "sheets") if k in old_schema}
    meta["labels"] = labels
    _refresh(conn, ds_id, df, meta)
    return lost


def delete(conn: sqlite3.Connection, ds_id: str) -> None:
    if not _row(conn, ds_id):
        raise DatasetError(f"没有这个数据集：{ds_id}")
    store.drop(ds_id)
    shutil.rmtree(store.dataset_dir(ds_id), ignore_errors=True)
    conn.execute("DELETE FROM datasets WHERE id = ?", (ds_id,))


def load(ds_id: str, columns: list[str] | None = None) -> pd.DataFrame:
    return store.read_df(ds_id, columns)

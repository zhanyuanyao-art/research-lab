"""Deck v2 (card 0196): one list of slides, drawn as an HTML preview inside the app and as a .pptx.

The look follows Leo's own course deck (IRT 国家间政治课程分享): a warm off-white page, deep-green titles, light cards
with a coloured top rule, rust and slate-blue accents, a small 「第几部分 · 部分名」 line above every title. Both renderers
lay a slide out on the same 1280 × 720 grid: the HTML in em units of a 16 px base that scales with the preview's width
(container query units), the pptx in inches at 96 px per inch, with font sizes at 0.75 pt per px.

Every string here may come from a paper, a plan or author code. The HTML template autoescapes; the pptx writer puts text
in runs, never in XML by hand; lengths are capped where a slide is assembled (report_content), and again here.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable

from .deck import FONT, Img, _page_number, _theme_fonts, clean

W_PX, H_PX = 1280, 720
PX = 13.333 / W_PX                # inches per px
PT = 0.75                         # points per px

PAPER, INK, GREEN, MUTED, LINE = "FBFAF6", "213831", "2B5148", "657267", "D5DBD4"
TINT, WARM, RUST, BLUE, OCHRE, CODE_BG, CODE_INK = "EEF2ED", "F2F0E8", "96513C", "365F77", "A58156", "213831", "E6ECE6"
TONES = {"green": (GREEN, TINT), "blue": (BLUE, TINT), "rust": (RUST, TINT), "warm": (OCHRE, WARM)}
HEAD_INK = {"green": GREEN, "blue": BLUE, "rust": RUST, "warm": "7D6140"}

MAX_SLIDES = 80
MAX_CARDS = 4
MAX_ROWS = 10
MAX_POINTS = 6
CODE_LINES = 9
CODE_CHARS = 88
BODY_PX, HEAD_PX, LINE_PX = 18, 20, 29          # card text, card heads, one body line (BODY_PX × 1.6)


@dataclass
class Card:
    head: str = ""
    body: str = ""
    tag: str = ""
    tone: str = "green"
    big: str = ""
    note: str = ""
    points: list[str] = field(default_factory=list)


@dataclass
class Slide:
    kind: str                     # cover | section | cards | flow | method | table | compare | chart | questions | points
    title: str = ""
    crumb: str = ""
    lede: str = ""
    src: str = ""
    # cover
    kicker: str = ""
    subtitle: str = ""
    meta: list[str] = field(default_factory=list)
    thumb: str = ""               # asset name
    # section
    no: str = ""
    sub: str = ""
    topics: list[str] = field(default_factory=list)
    # cards, flow, method, questions
    cards: list[Card] = field(default_factory=list)
    points: list[str] = field(default_factory=list)
    code: str = ""
    code_head: str = ""
    formula: str = ""
    # table and compare
    columns: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)
    hits: list[bool] = field(default_factory=list)       # compare: the row's reproduced value matches
    chips: list[str] = field(default_factory=list)       # table: the last column's tone per row, "" for plain text
    image: str = ""               # asset name: the paper's original, or a chart
    image_label: str = ""
    image2: str = ""              # compare: our picture of a figure, in place of the table
    image2_label: str = ""
    verdict: str = ""
    verdict_ok: bool = True
    verdict_note: str = ""


@dataclass
class Deck:
    title: str
    slides: list[Slide]
    assets: dict[str, Img] = field(default_factory=dict)   # name → picture, for both renderers


# ---------------------------------------------------------------- HTML

def render_html(deck: Deck, asset_url: Callable[[str], str], env: Any) -> str:
    """The slides as HTML for the preview page. `env` is the app's Jinja environment (autoescaping)."""
    tpl = env.get_template("partials/_deck_slides.html")
    sizes = {name: (img.w, img.h) for name, img in deck.assets.items()}
    return tpl.render(deck=deck, slides=deck.slides[:MAX_SLIDES], asset_url=asset_url, sizes=sizes,
                      code_lines=code_lines)


def code_lines(text: str) -> list[str]:
    lines = [ln[:CODE_CHARS] for ln in (text or "").split("\n")]
    return lines[:CODE_LINES] + (["…"] if len(lines) > CODE_LINES else [])


# ---------------------------------------------------------------- pptx primitives

def _in(px: float) -> int:
    from pptx.util import Inches
    return int(Inches(px * PX))


def _run_style(run: Any, px: float, color: str, bold: bool = False, mono: bool = False) -> None:
    from lxml import etree
    from pptx.dml.color import RGBColor
    from pptx.oxml.ns import qn
    from pptx.util import Pt
    f = run.font
    f.size, f.bold = Pt(px * PT), bold
    f.color.rgb = RGBColor.from_string(color)
    face = "Menlo" if mono else FONT
    f.name = face
    rpr = run._r.get_or_add_rPr()
    rpr.set("lang", "zh-CN")
    latin = rpr.find(qn("a:latin"))
    ea = etree.Element(qn("a:ea"))
    ea.set("typeface", FONT)
    latin.addnext(ea)


def _box(slide: Any, x: float, y: float, w: float, h: float, *, anchor: str = "t") -> Any:
    from pptx.enum.text import MSO_ANCHOR
    box = slide.shapes.add_textbox(_in(x), _in(y), _in(w), _in(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    return tf


def _para(tf: Any, first: bool) -> Any:
    return tf.paragraphs[0] if first else tf.add_paragraph()


def _write(slide: Any, x: float, y: float, w: float, h: float, text: str, px: float, color: str = INK, *,
           bold: bool = False, align: str = "l", anchor: str = "t", spacing: float = 1.0, mono: bool = False) -> None:
    from pptx.enum.text import PP_ALIGN
    tf = _box(slide, x, y, w, h, anchor=anchor)
    for i, line in enumerate((text or "").split("\n")):
        p = _para(tf, i == 0)
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        p.line_spacing = spacing
        r = p.add_run()
        _run_style(r, px, color, bold, mono)
        r.text = line


def _fill(slide: Any, x: float, y: float, w: float, h: float, color: str, line: str | None = None) -> Any:
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, _in(x), _in(y), _in(w), _in(h))
    s.shadow.inherit = False
    s.fill.solid()
    s.fill.fore_color.rgb = RGBColor.from_string(color)
    if line:
        s.line.color.rgb = RGBColor.from_string(line)
        s.line.width = 9525
    else:
        s.line.fill.background()
    return s


def _chip(slide: Any, x: float, y: float, text: str, color: str, px: float = 12) -> float:
    """A filled label; returns its width in px."""
    w = 16 + sum(px if ord(c) > 0x2E80 else px * 0.58 for c in text)
    s = _fill(slide, x, y, w, px + 10, color)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    _run_style(r, px, "FFFFFF")
    r.text = text
    return w


def _picture(slide: Any, img: Img, x: float, y: float, w: float, h: float, *, top: bool = True) -> None:
    k = min(w / img.w, h / img.h)
    iw, ih = img.w * k, img.h * k
    px_, py_ = x + (w - iw) / 2, y if top else y + (h - ih) / 2
    slide.shapes.add_picture(io.BytesIO(img.data), _in(px_), _in(py_), _in(iw), _in(ih))


def _lines(text: str, px: float, width: float) -> int:
    """Lines `text` takes at `px` in `width` px: a CJK character is as wide as the font size, Latin about half."""
    w = sum(px if ord(ch) > 0x2E80 else px * 0.55 for ch in text or "")
    return max(1, int(w // max(1.0, width)) + 1)


def _size_for(text: str, base: float, room: int) -> float:
    """A smaller font when the text would overflow: `room` is the characters that fit at `base`."""
    n = len(text or "")
    if n <= room:
        return base
    return max(base * 0.72, base * (room / n) ** 0.5)


# ---------------------------------------------------------------- pptx slides

def _chrome(slide: Any, s: Slide, n: int) -> None:
    from pptx.dml.color import RGBColor
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = RGBColor.from_string(PAPER)
    if s.kind in ("cover", "section"):
        if s.kind == "section":
            _page_number(slide, n)
        return
    if s.crumb:
        _write(slide, 64, 24, 1152, 20, s.crumb, 13, MUTED)
    _fill(slide, 64, 52, 1152, 1, LINE)
    _write(slide, 64, 66, 1152, 44, s.title, _size_for(s.title, 30, 34), INK, bold=True, anchor="m")
    if s.lede:
        _write(slide, 64, 114, 1152, 26, s.lede, _size_for(s.lede, 17, 66), MUTED)
    if s.src:
        _write(slide, 64, 690, 1000, 18, s.src, 12, MUTED)
    _page_number(slide, n)


def _card(slide: Any, c: Card, x: float, y: float, w: float, h: float) -> None:
    rule, bg = TONES.get(c.tone, TONES["green"])
    _fill(slide, x, y, w, h, bg)
    _fill(slide, x, y, w, 3, rule)
    cy = y + 20
    if c.tag:
        _chip(slide, x + 22, cy, c.tag, rule)
        cy += 32
    if c.big:
        _write(slide, x + 22, cy, w - 44, 40, c.big, _size_for(c.big, 34, max(4, int((w - 44) / 22))), rule, bold=True)
        cy += 46
    if c.head:
        lines = _lines(c.head, HEAD_PX, w - 44)
        _write(slide, x + 22, cy, w - 44, 30 * lines, c.head, HEAD_PX, HEAD_INK.get(c.tone, GREEN), bold=True)
        cy += 30 * lines + 8
    room = max(20, int((w - 44) / BODY_PX) * max(1, int((y + h - cy - 30) / LINE_PX)))
    body = c.body
    if c.points:
        body = "\n".join("· " + p for p in c.points[:MAX_POINTS])
    if body:
        bh = sum(LINE_PX * _lines(line, BODY_PX, w - 44) for line in body.split("\n"))
        _write(slide, x + 22, cy, w - 44, min(bh, y + h - cy - 20), body, _size_for(body, BODY_PX, room), INK, spacing=1.35)
        cy += bh
    if c.note:
        _write(slide, x + 22, min(cy + 10, y + h - 30), w - 44, 18, c.note, 12, MUTED)


def _card_height(c: Card, w: float) -> float:
    h = 42.0
    if c.tag:
        h += 32
    if c.big:
        h += 46
    if c.head:
        h += 30 * _lines(c.head, HEAD_PX, w - 44) + 8
    body = "\n".join("· " + p for p in c.points[:MAX_POINTS]) if c.points else c.body
    for line in (body or "").split("\n"):
        h += LINE_PX * _lines(line, BODY_PX, w - 44)
    if c.note:
        h += 30
    return h


def _cards_height(cards: list[Card], w: float, room: float) -> float:
    return min(room, max([220.0] + [_card_height(c, w) for c in cards]))


def _grid(n: int, x: float, w: float, gap: float = 18) -> list[tuple[float, float]]:
    cw = (w - gap * (n - 1)) / n
    return [(x + i * (cw + gap), cw) for i in range(n)]


def _code_height(text: str, w: float) -> float:
    lines = sum(_lines(ln, 14 * 0.62 / 0.55, w - 36) for ln in code_lines(text))
    return 22 + 28 + lines * 14 * 1.6


def _code(slide: Any, text: str, x: float, y: float, w: float, h: float, head: str = "") -> None:
    if head:
        _write(slide, x, y, w, 18, head, 12, MUTED)
        y += 22
        h -= 22
    _fill(slide, x, y, w, h, CODE_BG)
    lines = code_lines(text)
    _write(slide, x + 18, y + 14, w - 36, h - 28, "\n".join(lines), 14 if len(lines) <= 7 else 13, CODE_INK,
           mono=True, spacing=1.15)


def _table(slide: Any, columns: list[str], rows: list[list[str]], x: float, y: float, w: float, *,
           hits: list[bool] | None = None, chips: list[str] | None = None, row_h: float = 34, px: float = 14,
           widths: list[float] | None = None) -> None:
    from pptx.dml.color import RGBColor
    from pptx.util import Pt
    rows = rows[:MAX_ROWS * 2]
    n_rows, n_cols = len(rows) + 1, len(columns)
    if not n_cols:
        return
    shape = slide.shapes.add_table(n_rows, n_cols, _in(x), _in(y), _in(w), _in(row_h * n_rows))
    tbl = shape.table
    tbl.first_row = False
    tbl.horz_banding = False
    widths = widths or [1 / n_cols] * n_cols
    for i, f in enumerate(widths):
        tbl.columns[i].width = _in(w * f)
    for r in range(n_rows):
        tbl.rows[r].height = _in(row_h)
        for c in range(n_cols):
            cell = tbl.cell(r, c)
            cell.fill.background()
            cell.margin_left = cell.margin_right = _in(8)
            cell.margin_top = cell.margin_bottom = _in(4)
            text = columns[c] if r == 0 else (rows[r - 1][c] if c < len(rows[r - 1]) else "")
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            run = p.add_run()
            color = MUTED if r == 0 else INK
            if r and chips and c == n_cols - 1 and chips[r - 1]:
                color = "FFFFFF"
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string(GREEN if chips[r - 1] == "ok" else RUST)
            if r and hits and c == n_cols - 1 and hits[r - 1]:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string("DFE9E1")
            _run_style(run, 12 if r == 0 else px, color, bold=False)
            run.text = clean(text, 80)
            p.space_after = Pt(0)
    _fill(slide, x, y + row_h - 1, w, 1.5, GREEN if chips is not None else LINE)


def _draw(prs: Any, deck: Deck, s: Slide, n: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _chrome(slide, s, n)
    k = s.kind
    if k == "cover":
        if s.kicker:
            _write(slide, 80, 150, 700, 20, s.kicker, 14, MUTED)
        wide = s.thumb in deck.assets
        tw = 700 if wide else 1100
        _write(slide, 80, 184, tw, 160, s.title, _size_for(s.title, 44, 64 if wide else 100), INK, bold=True,
               anchor="b", spacing=1.05)
        if s.subtitle:
            _write(slide, 80, 356, tw, 34, s.subtitle, 24, GREEN)
        _fill(slide, 80, 420, 84, 3, GREEN)
        if s.meta:
            _write(slide, 80, 446, tw, 120, "\n".join(s.meta[:4]), 16, MUTED, spacing=1.4)
        if wide:
            img = deck.assets[s.thumb]
            _fill(slide, 880 - 1, 110 - 1, 302, 302 * img.h / img.w + 2, LINE)
            _picture(slide, img, 880, 110, 300, 300 * img.h / img.w)
        return
    if k == "section":
        _fill(slide, 0, 0, 14, H_PX, GREEN)
        _write(slide, 96, 200, 1000, 22, s.no, 15, MUTED)
        _write(slide, 96, 232, 1080, 72, s.title, _size_for(s.title, 56, 18), INK, bold=True)
        if s.sub:
            _write(slide, 96, 322, 1080, 32, s.sub, 22, GREEN)
        if s.topics:
            _write(slide, 96, 392, 1080, 24, " · ".join(s.topics[:6]), 15, MUTED)
        return
    top, bottom = (160 if s.lede else 132), 670
    if k in ("cards", "flow"):
        cards = s.cards[:MAX_CARDS]
        if k == "flow":
            arrow = 36
            cw = (1152 - arrow * (len(cards) - 1)) / max(1, len(cards))
            ch = _cards_height(cards, cw, bottom - top)
            for i, c in enumerate(cards):
                x = 64 + i * (cw + arrow)
                _card(slide, c, x, top, cw, ch)
                if i < len(cards) - 1:
                    _write(slide, x + cw, top + ch / 2 - 18, arrow, 36, "→", 26, GREEN, align="c")
        else:
            grid = _grid(max(1, len(cards)), 64, 1152)
            ch = _cards_height(cards, grid[0][1], bottom - top)
            for c, (x, cw) in zip(cards, grid):
                _card(slide, c, x, top, cw, ch)
        return
    if k == "method":
        lw = 560
        cy = top
        if s.cards:
            h = min(_card_height(s.cards[0], lw), 260)
            _card(slide, s.cards[0], 64, cy, lw, h)
            cy += h + 14
        if s.formula:
            _code(slide, s.formula, 64, cy, lw, min(bottom - cy, _code_height(s.formula, lw)), "模型")
        rx = 64 + lw + 32
        rw = 1216 - rx
        ry = top
        if len(s.cards) > 1:
            h = min(_card_height(s.cards[1], rw), 220)
            _card(slide, s.cards[1], rx, ry, rw, h)
            ry += h + 14
        if s.code:
            _code(slide, s.code, rx, ry, rw, min(bottom - ry, _code_height(s.code, rw)), s.code_head)
        return
    if k == "table":
        widths = [0.22, 0.26, 0.52] if len(s.columns) == 3 else None
        _table(slide, s.columns, s.rows[:MAX_ROWS], 64, top, 1152, chips=s.chips or None, widths=widths,
               row_h=50 if len(s.rows) <= 6 else 42, px=16)
        return
    if k == "compare":
        lw = 650
        if s.image in deck.assets:
            _write(slide, 64, top, lw, 18, s.image_label, 12, MUTED)
            _fill(slide, 64, top + 22, lw, bottom - top - 22, "FFFFFF", LINE)
            _picture(slide, deck.assets[s.image], 74, top + 32, lw - 20, bottom - top - 42)
        elif s.image_label:                       # the paper's picture was not captured: say so, never a blank half
            _write(slide, 64, top, lw, 18, s.image_label, 12, MUTED)
            _fill(slide, 64, top + 22, lw, bottom - top - 22, TINT)
        rx = 64 + lw + 24
        rw = 1216 - rx
        if s.image2 in deck.assets:
            _write(slide, rx, top, rw, 18, s.image2_label or "复现的图", 12, MUTED)
            _fill(slide, rx, top + 22, rw, bottom - top - 22, "FFFFFF", LINE)
            _picture(slide, deck.assets[s.image2], rx + 10, top + 32, rw - 20, bottom - top - 42)
            vy = bottom + 6
        elif s.columns and s.rows:
            _write(slide, rx, top, rw, 18, "复现结果 · 主要系数", 12, MUTED)
            rows = s.rows[:MAX_ROWS]
            rh = 30 if len(rows) > 7 else 34
            _table(slide, s.columns, rows, rx, top + 24, rw, hits=s.hits,
                   widths=[0.36, 0.24, 0.2, 0.2] if len(s.columns) == 4 else None, row_h=rh, px=13)
            vy = top + 24 + (len(rows) + 1) * rh + 16
        else:
            _fill(slide, rx, top + 22, rw, bottom - top - 22, TINT)
            vy = top + (bottom - top) / 2 - 12
        if s.verdict and not s.image2:
            cw = _chip(slide, rx + (0 if s.rows else 24), vy, s.verdict, GREEN if s.verdict_ok else RUST, 13)
            if s.verdict_note:
                _write(slide, rx + (0 if s.rows else 24) + cw + 10, vy + 1, rw - cw - 34, 22, s.verdict_note, 14, INK)
        return
    if k == "chart":
        if s.image in deck.assets:
            _picture(slide, deck.assets[s.image], 64, top, 1152, bottom - top)
        return
    if k == "questions":
        lw = 440
        if s.cards:
            _card(slide, s.cards[0], 64, top, lw, bottom - top)
        qx, qw = 64 + lw + 32, 1216 - (64 + lw + 32)
        qs = s.points[:5]
        step = min(96, (bottom - top) / max(1, len(qs)))
        for i, q in enumerate(qs):
            y = top + i * step
            _write(slide, qx, y + 6, 46, 30, f"{i + 1:02d}", 22, RUST, bold=True)
            _write(slide, qx + 54, y + 8, qw - 54, step - 14, q, _size_for(q, 18, 80), INK, spacing=1.3)
            _fill(slide, qx, y + step - 2, qw, 1, LINE)
        return
    if k == "points":
        body = "\n".join("· " + p for p in s.points[:MAX_POINTS * 2])
        _write(slide, 64, top, 1152, bottom - top, body, _size_for(body, 18, 900), INK, spacing=1.45)


def write_pptx(deck: Deck, out: Any) -> None:
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    _theme_fonts(prs)
    prs.core_properties.title = clean(deck.title, 200)
    prs.core_properties.author = "Research Lab"
    for n, s in enumerate(deck.slides[:MAX_SLIDES], 1):
        _draw(prs, deck, s, n)
    prs.save(out)


# ---------------------------------------------------------------- reading a method chapter into slides

_BOLD_ITEM = re.compile(r"^\d{1,2}\.[ \t]{1,4}\*\*([^*\n]{1,40})\*\*[ \t]{0,4}[：:][ \t]{0,4}([^\n]{0,400})", re.M)
_CODE_BLOCK = re.compile(r"```text\n(.{0,1200}?)\n```", re.S)


def first_sentence(text: str, limit: int = 90) -> str:
    text = re.sub(r"[*`]", "", (text or "").strip())
    for stop in ("。", "；"):
        i = text.find(stop)
        if 0 < i < limit:
            return text[:i + 1]
    return clean(text, limit)


def chapter_formula(md: str) -> str:
    m = _CODE_BLOCK.search(md or "")
    return m.group(1).strip() if m else ""


def chapter_assumptions(md: str, n: int = 5) -> list[tuple[str, str]]:
    return [(m.group(1).strip(), first_sentence(m.group(2), 70)) for m in _BOLD_ITEM.finditer(md or "")][:n]


def chapter_checks(md: str, n: int = 6) -> list[list[str]]:
    """The robustness table's rows: check, the assumption it tests, and the first sentence of how to read it."""
    rows = []
    for line in (md or "").split("\n"):
        if not line.startswith("|") or line.startswith("|---") or "检查哪条假设" in line:
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if len(cells) >= 3 and cells[0]:
            how = first_sentence(cells[2], 70)
            if len(how) < 24:                                  # 「和默认标准误并列。」 says little without the next sentence
                rest = cells[2][len(how):]
                how = clean(how + first_sentence(rest, 70 - len(how)), 80)
            rows.append([clean(cells[0], 24), clean(cells[1], 30), how])
        if len(rows) >= n:
            break
    return rows


def first_bullets(md: str, n: int = 3, limit: int = 70) -> list[str]:
    """The first `n` bullets of a section, each cut to its first sentence, without Markdown marks."""
    out = []
    for line in (md or "").split("\n"):
        if line.startswith("- "):
            out.append(first_sentence(line[2:], limit))
            if len(out) >= n:
                break
    return out


def first_paragraph(md: str, limit: int = 160) -> str:
    for block in (md or "").split("\n\n"):
        b = block.strip()
        if b and not b.startswith(("```", "|", "-", "#")):
            return first_sentence(b, limit) if len(b) > limit else re.sub(r"[*`]", "", b)
    return ""


def asset_file(folder: Path, name: str) -> Path | None:
    """A rendered asset by name, only directly inside `folder` and only plain names."""
    if not re.fullmatch(r"[A-Za-z0-9_\-]{1,80}\.(?:png|svg)", name or ""):
        return None
    p = folder / name
    return p if p.is_file() and not p.is_symlink() else None

"""展示稿 (card 0190, docs/design/REPLICATION-v4-DCL.md §3D): a .pptx for one replication plan, in the six parts of the DCL
lab's seminar deck. Rule based: no AI call, nothing is run, nothing is downloaded.

  封面          paper, authors, journal and year, 汇报人 (the owners on the checklist), the seminar date, "DCL sharing"
  01 研究问题与模型   from the paper's stored summary (papers.ai_summary_json) and the plan; 3 to 5 bullets a slide, a missing
                    field shows 「（待补充）」 and the slide carries 「待修改」
  02 数据收集与测量   the same, from the plan's data sources and variables
  03 结果       one slide per ticked 正文结果 item of the checklist (card 0187): left a crop of the paper, right our result
  04 稳健性检验   the same layout for the ticked appendix items
  05 结论与评价   結論 from the summary and a count of what was replicated; 评价 and 疑问 are empty templates for Leo

A result slide (the DCL deck's fixed layout). Left 「原始」: the caption and what follows it, cropped out of the paper's PDF
(the caption is found with search_for on a page that is read upright, card 0182's rule; the crop runs to the next gap in the
page or to its bottom; when no caption is found the whole page, small). Right 「现在」: a table from the latest valid run's
results.json as a pptx table, its cells tinted by compare.compare against the plan's targets; or the run's figure (a PNG, or
the first page of a PDF); or 「尚未复现」 / 「等数据」 (card 0188's waiting scripts). Footer: 「来源：正文 表 3（论文第 7 页）；
meetingAnalysis.R　判定：完全一致 3 · 对不上 1」 and the page number.

The deck never carries a judgment: 评价 and 疑问 stay empty, and the conclusion only restates the paper's own summary and counts
what the run found.

Outside text. Captions, summary fields, plan rows, script names and the paper's PDF are somebody else's text (AGENTS.md,
"Scanning outside text"): every string is cut before any pattern sees it, control and bidirectional characters are dropped,
counts are capped (slides, table cells, layout blocks, lines kept) and the rest is counted, a page lookup indexes labels once and
looks hits up by position, and past DECK_SECONDS the remaining crops are skipped. The route runs make() in the thread pool.

Output: <plan workspace>/exports/<plan>-deck.pptx, written by a temporary file and a rename, never through a link."""
from __future__ import annotations

import io
import json
import logging
import os
import posixpath
import re
import sqlite3
import stat
import subprocess
import sys
import tempfile
import time
import uuid
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, NamedTuple

from ..config import get_paths
from . import compare, coverage, plans
from . import workspace as W

log = logging.getLogger("research_lab.replication.deck")

MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
EXPORTS = "exports"
FONT = "PingFang SC"
FALLBACK_FONTS = {"Hans": "Microsoft YaHei", "Hant": "Microsoft JhengHei"}      # the theme's per-script fallback

# Ceilings
MAX_SLIDES = 60
MAX_BULLETS = 5
BULLET_CHARS = 200
TITLE_CHARS = 160
SHORT_CHARS = 40
FIELD_CHARS = 600
SUMMARY_BYTES = 200_000
MAX_LIST = 40                         # plan rows read from a list
MAX_VARIABLE_SLIDES = 3
VARIABLES_PER_SLIDE = 5
MAX_TABLE_COLS = 8
MAX_TABLE_ROWS = 13
MAX_RESULT_ROWS = 5_000               # result rows read for one table
MAX_COMPARED = 2_000                  # compare items read
MAX_RESULT_TERMS = 300
MAX_RESULT_MODELS = 60
TERM_CHARS = 44
MAX_FILE_NAMES = 3
MAX_WAITING = 200
MAX_PAGES_LABELLED = 3_000
MAX_CANDIDATE_PAGES = 4
MAX_LAYOUT_LINES = 6_000
MAX_LAYOUT_BLOCKS = 3_000
MAX_IMAGE_RECTS = 400
MAX_DRAWINGS = 4_000
MAX_LINE = 500                        # characters of a page line the caption rules read
HEAD_CHARS = 600
MAX_HITS = 24
MAX_OUTPUT_FILES = 3_000
IMAGE_MAX_PX = 1_800
ROW_IN = 0.28                         # a table row on a slide: two lines of CELL_PT
CELL_PT = 7
IMAGE_MAX_PIXELS = 16_000_000
IMAGE_MAX_BYTES = 12 * 1024 * 1024
IMAGE_TOTAL_BYTES = 150 * 1024 * 1024   # all the pictures of one deck; past it a slide says 「图过大，未放入」
PDF_FIGURE_BYTES = 20 * 1024 * 1024     # a run's PDF figure larger than this is not rendered
RENDER_SECONDS = 12.0                   # one PDF figure, in a child process that is killed at the limit
CROP_DPI = 150
PAGE_DPI = 100
GAP_PT = 20.0                         # a gap this wide in the page ends the crop …
FIRST_GAP_PT = 48.0                   # … except between the caption and what it captions (a table sets its head apart)
DECK_SECONDS = 120.0

INK, NAVY, MUTED, RULE, PANEL = "1F2937", "1E3A5F", "6B7280", "D1D5DB", "F3F4F6"
TAG_FILL, TAG_TEXT, HEAD_FILL = "FEF3C7", "92400E", "E8EEF7"
VERDICT_FILL = {"exact": "DCEFE0", "within": "DCEFE0", "direction": "FBEBC8", "mismatch": "F6D7D5"}
VERDICT_NAMES = {"exact": "完全一致", "within": "在误差范围内", "direction": "方向一致", "mismatch": "对不上", "missing": "没找到对应输出"}
PARTS = (("01", "研究问题与模型"), ("02", "数据收集与测量"), ("03", "结果"), ("04", "稳健性检验"), ("05", "结论与评价"))
MISSING = "（待补充）"


class DeckError(ValueError):
    """Refused, in Chinese, as the page shows it."""


# ---------------------------------------------------------------- text from outside

_BAD = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f\ud800-\udfff﷐-﷯￾￿​-‏"
                  " -‮⁦-⁩]")


def clean(text: Any, limit: int = BULLET_CHARS) -> str:
    """One line, no control or bidirectional characters, at most `limit` characters. Cut before anything else reads it."""
    s = "" if text is None else (text if isinstance(text, str) else str(text) if isinstance(text, (int, float)) else "")
    s = " ".join(_BAD.sub(" ", s[: limit * 3 + 60]).split())
    return s if len(s) <= limit else s[: max(0, limit - 1)].rstrip() + "…"


def _nt(v: Any) -> str:
    return compare.norm_table(clean(v, 60))


def _nm(v: Any) -> str:
    return compare.norm_model(clean(v, 40))


def _nterm(v: Any) -> str:
    return compare.norm_term(clean(v, 80))


def _list(v: Any, n: int = MAX_LIST) -> list[Any]:
    return list(v[:n]) if isinstance(v, list) else []


def _dict(v: Any) -> dict[str, Any]:
    return v if isinstance(v, dict) else {}


# ---------------------------------------------------------------- the model of a deck

@dataclass
class Bullet:
    text: str
    missing: bool = False


@dataclass
class Img:
    data: bytes
    w: int
    h: int


@dataclass
class Cell:
    lines: list[str]
    fill: str | None = None
    bold: bool = False


@dataclass
class Matrix:
    header: list[str]
    rows: list[list[Cell]]
    notes: list[str] = field(default_factory=list)
    legend: str = ""


@dataclass
class Spec:
    kind: str                                   # cover | part | text | result | blank
    title: str = ""
    crumb: str = ""
    bullets: list[Bullet] = field(default_factory=list)
    thin: bool = False
    tag: str = ""                               # 「待修改」 / 「待填写」
    footer: str = ""
    number: str = ""                            # part number
    cover: dict[str, str] = field(default_factory=dict)
    owner: str = ""
    left: Img | None = None
    left_note: str = ""
    right: str = ""                             # table | image | note
    right_img: Img | None = None
    matrix: Matrix | None = None
    right_note: str = ""
    right_detail: str = ""


@dataclass
class DeckResult:
    path: Path
    slides: int
    notes: list[str]
    replicated: int = 0
    pending: int = 0


# ---------------------------------------------------------------- what the plan holds

@dataclass
class Ctx:
    plan: dict[str, Any]
    paper: dict[str, Any]
    summary: dict[str, Any]
    rows: list[dict[str, Any]]
    stored: bool
    run: dict[str, Any] | None = None
    comparison: dict[str, Any] | None = None
    reviewed: bool = False
    waiting: dict[str, list[str]] = field(default_factory=dict)
    verdicts: dict[str, dict[tuple[str, str], str]] = field(default_factory=dict)
    counts: dict[str, dict[str, int]] = field(default_factory=dict)
    seminar_date: str = ""
    notes: list[str] = field(default_factory=list)

    @property
    def plan_id(self) -> str:
        return str(self.plan["id"])

    @property
    def tables(self) -> list[dict[str, Any]]:
        if self.run is None:
            return []
        results = _dict(_dict(self.run.get("result")).get("results"))
        return [t for t in _list(results.get("tables"), 400) if isinstance(t, dict)]


def _row_of(it: coverage.Item) -> dict[str, Any]:
    return {"id": None, "label": it.label, "caption": it.caption, "kind": it.kind, "page": it.page, "source": it.source,
            "group": it.group, "scripts": it.scripts, "outputs": it.outputs, "ticked": it.group == "main", "owner": ""}


def _summary_of(raw: Any) -> dict[str, Any]:
    if not isinstance(raw, str) or not raw or len(raw) > SUMMARY_BYTES:
        return {}
    try:
        s = json.loads(raw)
    except ValueError:
        return {}
    if isinstance(s, dict) and isinstance(s.get("summary"), dict):
        s = s["summary"]
    return s if isinstance(s, dict) else {}


def gather(conn: sqlite3.Connection, plan_id: str, *, seminar_date: str = "") -> Ctx:
    plan = plans.get(conn, plan_id)
    if plan is None:
        raise DeckError(f"没有这个复现方案：{plan_id}")
    paper: dict[str, Any] = {}
    if plan.get("paper_id"):
        r = conn.execute("SELECT title, year, venue, authors_json, ai_summary_json FROM papers WHERE id = ?", (plan["paper_id"],)).fetchone()
        if r is not None:
            paper = dict(r)
    notes: list[str] = []
    rows = coverage.items(conn, plan_id)
    stored = bool(rows)
    if not stored:
        try:
            result = coverage.scan(conn, plan_id)
            rows = [_row_of(i) for i in result.items]
            notes += result.notes
        except (coverage.CoverageError, plans.PlanError, OSError, ValueError) as e:
            notes.append(f"清单读不出来：{clean(str(e), 120)}")
        notes.append("全书复现清单还没有 Scan，按默认勾选（正文结果）出稿。")
    run = W.latest_valid(conn, plan_id)
    ctx = Ctx(plan, paper, _summary_of(paper.get("ai_summary_json")), rows, stored, run=run, seminar_date=clean(seminar_date, SHORT_CHARS),
              notes=notes, reviewed=plans.targets_reviewed(plan))
    if run is None:
        ctx.notes.append("还没有可用的复现运行：结果页写「尚未复现」。")
    else:
        results = _dict(_dict(run.get("result")).get("results"))
        targets = plans.target_records(conn, plan_id)
        rows_n = sum(len(_list(_dict(t).get("rows"), MAX_RESULT_ROWS + 1)) for t in _list(results.get("tables"), 400))
        if rows_n > MAX_RESULT_ROWS:
            ctx.notes.append("复现结果的行数太多，没有和论文报告值对照。")
        elif targets and results:
            try:
                ctx.comparison = compare.compare(targets, results, _dict(run["result"].get("mapping")) or None)
                ctx.verdicts, ctx.counts = index_comparison(ctx.comparison)
            except (ValueError, TypeError, KeyError, ArithmeticError) as e:
                ctx.notes.append(f"和论文报告值对不上号，判定留空：{clean(str(e), 80)}")
    try:
        from . import missing
        ctx.waiting = missing.waiting_scripts(conn, plan_id)
    except (ValueError, OSError, sqlite3.Error) as e:
        log.info("waiting scripts unavailable: %s", e)
    return ctx


# ---------------------------------------------------------------- the paper's page: where a caption is, what follows it

class Block(NamedTuple):
    x0: float
    y0: float
    x1: float
    y1: float
    lines: int
    head: str


class Piece(NamedTuple):
    x0: float
    y0: float
    x1: float
    y1: float
    stop: bool = False


class Layout(NamedTuple):
    lines: list[tuple[float, float, float, float, str]]
    blocks: list[Block]
    images: list[tuple[float, float, float, float]]
    width: float                                # the page as it reads (reading frame), not as it lies
    height: float
    disp: tuple[float, float]                   # the page as displayed (rotation applied)
    turn: int                                   # degrees clockwise the displayed page must turn to read left to right
    rot: Any                                    # page.rotation_matrix: raw coordinates -> displayed


_CAPTION_START = re.compile(r"^(?:appendix[ ]+)?(?:table|tab\.?|figure|fig\.?|表|图)[ ]*[a-z]?\.?[ ]?\d", re.I)
SIDEWAYS_SHARE = 0.6                             # as intake.extract: the share of characters that run the same way
SIDEWAYS_MIN_CHARS = 80


def _norm_line(text: str) -> str:
    return " ".join(text[:MAX_LINE].split()).lower()


def _alnum(text: str) -> str:
    return "".join(c for c in text[:MAX_LINE * 3].lower() if c.isalnum())


def starts_with_label(line: str, needle: str) -> bool:
    """The line opens with the label and the label ends there ("Table 3" is not "Table 30" or "Table 3a")."""
    s, n = _norm_line(line), needle.lower()
    return s.startswith(n) and (len(s) == len(n) or not s[len(n)].isalnum())


def needles_for(kind: str, label: str) -> list[str]:
    num = clean(label, 20).rsplit(" ", 1)[-1]
    if label[:1] in ("表", "图"):
        word = "表" if kind == "table" else "图"
        return [f"{word}{num}", f"{word} {num}"]
    return [f"{w} {num}" for w in (("Table", "Tab.") if kind == "table" else ("Figure", "Fig."))]


def is_prose(b: Block, page_w: float) -> bool:
    """A paragraph of running text (a wall for the crop): four lines or more, wide, nearly no digits, not a table note."""
    if b.lines < 4 or (b.x1 - b.x0) < 0.5 * page_w or len(b.head) < 200:
        return False
    if b.head[:12].lower().lstrip(" *†‡§.:").startswith(("note", "source")):
        return False
    return sum(1 for c in b.head if c.isdigit()) <= 0.03 * len(b.head)


def to_reading(box: tuple[float, float, float, float], rot: Any, disp: tuple[float, float], turn: int) -> tuple[float, float, float, float]:
    """A box in the page's own coordinates (what get_text and search_for give) in the frame the page reads in: the page's
    /Rotate first, then the quarter turns that make its text run left to right (a landscape table set sideways)."""
    x0, y0, x1, y1 = box
    pts = [(rot.a * x + rot.c * y + rot.e, rot.b * x + rot.d * y + rot.f) for x in (x0, x1) for y in (y0, y1)]
    w, h = disp
    if turn == 90:
        pts = [(h - y, x) for x, y in pts]
    elif turn == 270:
        pts = [(y, w - x) for x, y in pts]
    elif turn == 180:
        pts = [(w - x, h - y) for x, y in pts]
    xs, ys = [p[0] for p in pts], [p[1] for p in pts]
    return min(xs), min(ys), max(xs), max(ys)


def to_displayed(box: tuple[float, float, float, float], disp: tuple[float, float], turn: int) -> tuple[float, float, float, float]:
    """The inverse of the quarter turns in to_reading: a box of the reading frame in the displayed page's coordinates."""
    rx0, ry0, rx1, ry1 = box
    w, h = disp
    if turn == 90:
        return ry0, h - rx1, ry1, h - rx0
    if turn == 270:
        return w - ry1, rx0, w - ry0, rx1
    if turn == 180:
        return w - rx1, h - ry1, w - rx0, h - ry0
    return box


def reading_turn(dirs: list[tuple[float, float, int]], rot: Any) -> int:
    """Quarter turns (clockwise, 0, 90, 180 or 270) of the displayed page after which most of its characters run left to
    right: the line directions are in the page's own coordinates, so /Rotate is applied to them first (card 0182's
    intake.extract.sideways_turn leaves it out, which is right for text and wrong for a picture)."""
    chars = {0: 0, 90: 0, 180: 0, 270: 0}
    for dx, dy, n in dirs:
        x, y = rot.a * dx + rot.c * dy, rot.b * dx + rot.d * dy
        turn = 0 if x >= 0.95 else 90 if y <= -0.95 else 180 if x <= -0.95 else 270 if y >= 0.95 else None
        if turn is not None:
            chars[turn] += n
    total = sum(chars.values())
    best = max(chars, key=lambda k: chars[k])
    return best if best and total >= SIDEWAYS_MIN_CHARS and chars[best] >= SIDEWAYS_SHARE * total else 0


def layout_of(page: Any) -> Layout:
    """Lines, text blocks and image boxes of a page in its reading frame, every list capped."""
    d = page.get_text("dict", flags=0)
    rot = page.rotation_matrix
    disp = (page.rect.width, page.rect.height)
    raw_lines: list[tuple[tuple[float, float, float, float], str]] = []
    raw_blocks: list[tuple[tuple[float, float, float, float], int, str]] = []
    dirs: list[tuple[float, float, int]] = []
    for b in _list(d.get("blocks"), MAX_LAYOUT_BLOCKS * 4):
        if not isinstance(b, dict) or b.get("type") != 0 or len(raw_blocks) >= MAX_LAYOUT_BLOCKS:
            continue
        head: list[str] = []
        used = n = 0
        for ln in _list(b.get("lines"), 400):
            if len(raw_lines) >= MAX_LAYOUT_LINES:
                break
            parts, got = [], 0
            for sp in _list(ln.get("spans"), 200):
                t = sp.get("text") or ""
                if got < MAX_LINE:
                    parts.append(t[: MAX_LINE - got])
                got += len(t)
            text = "".join(parts)
            dx, dy = ln.get("dir") or (1.0, 0.0)
            dirs.append((dx, dy, sum(1 for c in text if not c.isspace())))
            raw_lines.append((tuple(ln["bbox"]), text))        # type: ignore[arg-type]
            n += 1
            if used < HEAD_CHARS:
                head.append(text)
                used += len(text) + 1
        raw_blocks.append((tuple(b["bbox"]), n, " ".join(head)[:HEAD_CHARS]))      # type: ignore[arg-type]
    turn = reading_turn(dirs, rot)
    lines = [(*to_reading(box, rot, disp, turn), text) for box, text in raw_lines]
    blocks = [Block(*to_reading(box, rot, disp, turn), n, head) for box, n, head in raw_blocks]
    images: list[tuple[float, float, float, float]] = []
    try:
        for info in _list(page.get_image_info(), MAX_IMAGE_RECTS):
            images.append(to_reading(tuple(info["bbox"]), rot, disp, turn))      # type: ignore[arg-type]
    except (RuntimeError, ValueError, KeyError):
        pass
    w, h = disp if turn % 180 == 0 else (disp[1], disp[0])
    return Layout(lines, blocks, images, w, h, disp, turn, rot)


def find_caption(page: Any, lay: Layout, needles: list[str], caption: str) -> tuple[int, int] | None:
    """(index of the caption's line, tier). Tier 1: the line opens with the label and the caption's first words follow in
    this line or the next two; tier 2: only the label opens the line. A hit in the middle of a sentence is not a caption."""
    head = _alnum(caption)[:24]
    found: dict[int, int] = {}
    seen = 0
    for needle in needles:
        for r in page.search_for(needle):
            seen += 1
            if seen > MAX_HITS:
                break
            hx0, hy0, hx1, hy1 = to_reading((r.x0, r.y0, r.x1, r.y1), lay.rot, lay.disp, lay.turn)
            cx, cy = (hx0 + hx1) / 2, (hy0 + hy1) / 2
            for i, (x0, y0, x1, y1, text) in enumerate(lay.lines):
                if x0 - 1 <= cx <= x1 + 1 and y0 - 1 <= cy <= y1 + 1:
                    if starts_with_label(text, needle) and i not in found:
                        nxt = " ".join(t[4] for t in lay.lines[i:i + 3])
                        found[i] = 1 if not head or head in _alnum(nxt) else 2
                    break
    if not found:
        return None
    best = min(found, key=lambda i: (found[i], lay.lines[i][1]))
    return best, found[best]


def _overlap(p: Piece | Block, x0: float, x1: float) -> bool:
    w = min(p.x1 - p.x0, x1 - x0)
    return w > 0 and (min(p.x1, x1) - max(p.x0, x0)) >= 0.25 * w


def crop_box(cap: tuple[float, float, float, float], pieces: list[Piece], page: tuple[float, float], kind: str,
             gap: float = GAP_PT, first_gap: float = FIRST_GAP_PT) -> tuple[float, float, float, float]:
    """From the caption down to the next gap in the page (a wall, a new caption or a paragraph, ends it early); a figure whose
    caption sits below it is cropped upward instead. Sorting is the only superlinear step; pieces are capped by the caller."""
    ux0, uy0, ux1, bottom = cap
    top = uy0
    started = False
    for p in sorted((p for p in pieces if p.y0 > uy0 + 2), key=lambda p: (p.y0, p.x0)):
        if not _overlap(p, ux0, ux1):
            continue
        if p.y0 - bottom > (gap if started else first_gap) or p.stop:
            break
        bottom, ux0, ux1, started = max(bottom, p.y1), min(ux0, p.x0), max(ux1, p.x1), True
    if kind == "figure" and not started:
        for p in sorted((p for p in pieces if p.y1 <= uy0 + 2), key=lambda p: -p.y1):
            if not _overlap(p, ux0, ux1):
                continue
            if top - p.y1 > (gap if started else first_gap) or p.stop:
                break
            top, ux0, ux1, started = min(top, p.y0), min(ux0, p.x0), max(ux1, p.x1), True
    return ux0, top, ux1, bottom


def _drawings(page: Any, lay: Layout) -> list[tuple[float, float, float, float]]:
    out: list[tuple[float, float, float, float]] = []
    try:
        for d in page.get_drawings()[:MAX_DRAWINGS]:
            r = d["rect"]
            if r.width > 0.5 or r.height > 0.5:
                out.append(to_reading((r.x0, r.y0, r.x1, r.y1), lay.rot, lay.disp, lay.turn))
    except (RuntimeError, ValueError, KeyError):
        pass
    return out


def render(page: Any, clip: Any | None, dpi: int = CROP_DPI, turn: int = 0) -> Img:
    """A picture of the page (or of `clip`, in displayed coordinates), turned `turn` degrees clockwise."""
    import pymupdf
    rect = clip if clip is not None else page.rect
    scale = min(dpi / 72.0, IMAGE_MAX_PX / max(rect.width, rect.height, 1.0))
    pix = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), clip=clip, alpha=False)
    data, w, h = pix.tobytes("png"), pix.width, pix.height
    if turn:
        from PIL import Image
        im = Image.open(io.BytesIO(data)).rotate(-turn, expand=True)
        buf = io.BytesIO()
        im.save(buf, "PNG")
        data, w, h = buf.getvalue(), im.width, im.height
    return Img(data, w, h)


@dataclass
class Crop:
    img: Img
    mode: str                                   # caption | page
    page: int


def _try_page(page: Any, lay: Layout, kind: str, label: str, caption: str) -> tuple[int, Img] | None:
    import pymupdf
    hit = find_caption(page, lay, needles_for(kind, label), caption)
    if hit is None:
        return None
    i, tier = hit
    x0, y0, x1, y1, _ = lay.lines[i]
    cap_block = next((b for b in lay.blocks if b.x0 - 1 <= (x0 + x1) / 2 <= b.x1 + 1 and b.y0 - 2 <= y0 <= b.y1 + 1), None)
    bx0, bx1, by1 = (cap_block.x0, cap_block.x1, max(cap_block.y1, y1)) if cap_block else (x0, x1, y1)
    cap = (min(bx0, x0), y0, max(bx1, x1), by1)
    pieces = []
    for b in lay.blocks:
        if cap_block is not None and (b.x0, b.y0, b.x1, b.y1) == (cap_block.x0, cap_block.y0, cap_block.x1, cap_block.y1):
            continue
        pieces.append(Piece(b.x0, b.y0, b.x1, b.y1, bool(_CAPTION_START.match(b.head[:24])) or is_prose(b, lay.width)))
    pieces += [Piece(*r) for r in lay.images]
    if kind == "figure":
        pieces += [Piece(*r) for r in _drawings(page, lay)]
    x0, top, x1, bottom = crop_box(cap, pieces, (lay.width, lay.height), kind)
    shown = to_displayed((x0 - 6, top - 6, x1 + 6, bottom + 6), lay.disp, lay.turn)
    clip = pymupdf.Rect(*shown) & page.rect
    if clip.is_empty or clip.height < 24 or clip.width < 24:
        return None
    return tier, render(page, clip, turn=lay.turn)


def crop_original(doc: Any, labels: list[str], row: dict[str, Any]) -> Crop | None:
    """The caption of `row` and what follows it, cut out of the paper's PDF; the whole page (small) when no caption is
    found; None when the page is not in the file. A page that lies sideways is read, and cropped, as it reads."""
    page_label = clean(row.get("page"), 12)
    cands = [i for i, l in enumerate(labels) if l == page_label][:MAX_CANDIDATE_PAGES]
    if not cands and page_label.isdigit() and 1 <= int(page_label) <= doc.page_count:
        cands = [int(page_label) - 1]
    kind, label, caption = str(row.get("kind") or "table"), clean(row.get("label"), 24), clean(row.get("caption"), 240)
    second: Crop | None = None
    first_lay: Layout | None = None
    for i in cands:
        page = doc[i]
        lay = layout_of(page)
        first_lay = first_lay or lay
        got = _try_page(page, lay, kind, label, caption)
        if got and got[0] == 1:
            return Crop(got[1], "caption", i)
        if got and second is None:
            second = Crop(got[1], "caption", i)
    if second is not None:
        return second
    if cands and first_lay is not None:
        return Crop(render(doc[cands[0]], None, PAGE_DPI, first_lay.turn), "page", cands[0])
    return None


class PaperPdf:
    """The paper's PDF files, opened once each for the whole build and closed after it."""

    def __init__(self, conn: sqlite3.Connection, paper_id: str | None):
        self.files: list[tuple[Any, list[str]]] = []
        self.notes: list[str] = []
        if not paper_id:
            return
        import pymupdf
        from ..intake import extract
        root = os.path.realpath(get_paths().data)
        for f in conn.execute("SELECT rel_path FROM paper_files WHERE paper_id = ? ORDER BY is_primary DESC, id", (paper_id,)):
            path = os.path.realpath(get_paths().data / str(f["rel_path"]))
            try:
                if not path.startswith(root + os.sep) or not stat.S_ISREG(os.lstat(path).st_mode):
                    raise OSError("not a regular file inside the data folder")
                extract.check_pdf(Path(path))
                doc = pymupdf.open(path)
                labels = [extract._label(doc[i]) or str(i + 1) for i in range(min(doc.page_count, MAX_PAGES_LABELLED))]
                self.files.append((doc, labels))
            except (OSError, RuntimeError, ValueError, extract.RejectedFile) as e:
                self.notes.append(f"论文 PDF 打不开，没有截原图：{clean(str(e), 80)}")
                self.files.append((None, []))

    def close(self) -> None:
        for doc, _ in self.files:
            if doc is not None:
                doc.close()
        self.files = []

    def crop(self, row: dict[str, Any]) -> Crop | None:
        """The primary file serves the main text, every other file the appendix (the order coverage.sources_of reads them in)."""
        mine = self.files[:1] if row.get("source") == "main" else self.files[1:]
        for doc, labels in mine:
            if doc is None:
                continue
            got = crop_original(doc, labels, row)
            if got is not None:
                return got
        return None


# ---------------------------------------------------------------- our result: a table or a figure

def fmt_number(x: float) -> str:
    if x != x or x in (float("inf"), float("-inf")):
        return ""
    if abs(x) >= 1000 or (x == round(x) and abs(x) >= 10):
        text = f"{x:,.0f}"
    elif x != 0 and abs(x) < 0.001:
        text = f"{x:.1e}"
    else:
        text = f"{x:.3f}"
    return text.replace("-", "−")


def _num(v: Any) -> float | None:
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def cell_of(row: dict[str, Any], fill: str | None) -> Cell:
    est, se, p = _num(row.get("estimate")), _num(row.get("se")), _num(row.get("p"))
    stars = "*" * compare.p_band(p) if p is not None and 0 <= p <= 1 else ""
    first = (fmt_number(est) + stars) if est is not None else "—"
    second = ""
    if se is not None:
        second = f"({fmt_number(se)})"
    else:
        ci = row.get("ci")
        if isinstance(ci, list) and len(ci) == 2 and all(_num(x) is not None for x in ci):
            second = f"[{fmt_number(ci[0])}, {fmt_number(ci[1])}]"
    return Cell([first, second] if second else [first], fill)


def index_comparison(comparison: dict[str, Any] | None) -> tuple[dict[str, dict[tuple[str, str], str]], dict[str, dict[str, int]]]:
    """compare.compare's items by table, one pass: (table → (model, term) of our result → verdict, table → verdict counts).
    Every label is normalised once, so a deck of sixty slides does not walk the items sixty times."""
    cells: dict[str, dict[tuple[str, str], str]] = {}
    counts: dict[str, dict[str, int]] = {}
    for it in _list(_dict(comparison).get("items"), MAX_COMPARED):
        it = _dict(it)
        verdict = it.get("verdict")
        if verdict not in VERDICT_NAMES:
            continue
        table = _nt(it.get("table"))
        c = counts.setdefault(table, {})
        c[verdict] = c.get(verdict, 0) + 1
        m = it.get("matched")
        if isinstance(m, dict) and verdict in VERDICT_FILL:
            cells.setdefault(_nt(m.get("table")), {})[(_nm(m.get("model")), _nterm(m.get("term")))] = verdict
    return cells, counts


def matrix_of(tables: list[dict[str, Any]], label: str, verdicts: dict[tuple[str, str], str] | None = None) -> Matrix | None:
    """The models as columns and the terms as rows, two lines a cell (estimate with stars, then the standard error), tinted
    where a target of the plan was matched. At most MAX_TABLE_COLS models and MAX_TABLE_ROWS terms are drawn."""
    want = _nt(label)
    mine = [t for t in tables if _nt(t.get("label")) == want and _list(t.get("rows"), 1)]
    if not mine:
        return None
    verdicts = verdicts or {}
    models: dict[str, int] = {}
    terms: dict[str, int] = {}
    cells: dict[tuple[int, int], Cell] = {}
    read = 0
    for t in mine:
        model = clean(t.get("model"), 16) or "—"
        if model not in models:
            if len(models) >= MAX_RESULT_MODELS:
                continue
            models[model] = len(models)
        for r in _list(t.get("rows"), MAX_RESULT_ROWS):
            if not isinstance(r, dict) or read >= MAX_RESULT_ROWS:
                continue
            read += 1
            term = clean(r.get("term"), TERM_CHARS)
            if term not in terms:
                if len(terms) >= MAX_RESULT_TERMS:
                    continue
                terms[term] = len(terms)
            v = verdicts.get((_nm(t.get("model")), _nterm(r.get("term"))))
            cells[(terms[term], models[model])] = cell_of(r, VERDICT_FILL.get(v or ""))
    cols = min(len(models), MAX_TABLE_COLS)
    rows_n = min(len(terms), MAX_TABLE_ROWS)
    names = list(terms)
    body = [[Cell([names[i]], bold=False), *[cells.get((i, j), Cell([""])) for j in range(cols)]] for i in range(rows_n)]
    notes = []
    if len(models) > cols or len(terms) > rows_n:
        notes.append("另有" + (f" {len(models) - cols} 个模型" if len(models) > cols else "")
                     + ("、" if len(models) > cols and len(terms) > rows_n else "")
                     + (f" {len(terms) - rows_n} 行" if len(terms) > rows_n else "") + "未列出")
    legend = "绿：一致　黄：方向一致　红：对不上" if any(c.fill for r in body for c in r) else ""
    return Matrix(["", *list(models)[:cols]], body, notes, legend)


def _png_bounded(data: bytes) -> Img | None:
    from PIL import Image
    try:
        im = Image.open(io.BytesIO(data))
        if im.width * im.height > IMAGE_MAX_PIXELS:
            return None
        im.load()
    except (OSError, ValueError, SyntaxError, Image.DecompressionBombError):
        return None
    im.thumbnail((IMAGE_MAX_PX, IMAGE_MAX_PX))
    buf = io.BytesIO()
    (im if im.mode in ("RGB", "RGBA", "L") else im.convert("RGB")).save(buf, "PNG")
    if buf.tell() > IMAGE_MAX_BYTES:
        buf = io.BytesIO()
        im.convert("RGB").save(buf, "JPEG", quality=80)
    return Img(buf.getvalue(), im.width, im.height)


def _read_file(path: Path, cap: int = IMAGE_MAX_BYTES * 2) -> bytes | None:
    """A plain single-link file of at most `cap` bytes, read through a descriptor opened without following a link."""
    try:
        fd = os.open(path, os.O_RDONLY | os.O_NOFOLLOW | os.O_NONBLOCK)
    except OSError:
        return None
    try:
        st = os.fstat(fd)
        if not stat.S_ISREG(st.st_mode) or st.st_nlink > 1 or st.st_size > cap:
            return None
        with os.fdopen(fd, "rb", closefd=False) as fh:
            return fh.read(cap)
    except OSError:
        return None
    finally:
        os.close(fd)


def row_key(row: dict[str, Any]) -> str:
    return coverage.mkey(str(row.get("kind") or "table"), clean(row.get("label"), 24).rsplit(" ", 1)[-1])


class Figures:
    """The pictures a run wrote (outputs/, plain files only), indexed once by file name and by the label the name carries
    (Figure_7.png), so each figure slide is two dict lookups."""

    RANK = {".png": 0, ".jpg": 1, ".jpeg": 1, ".pdf": 2}         # a picture is cheaper than a PDF: it wins when a run wrote both

    def __init__(self, plan_id: str | None = None, run_id: str | None = None):
        self.by_stem: dict[str, Path] = {}
        self.by_key: dict[str, Path] = {}
        if not plan_id or not run_id:
            return
        root = W.run_dir(plan_id, run_id) / "outputs"
        if not root.is_dir() or root.is_symlink():
            return
        files = [f for f in W.regular_files(root)[:MAX_OUTPUT_FILES]
                 if f.suffix.lower() in self.RANK and not f.name.lower().startswith("rplots")]
        for f in sorted(files, key=lambda f: (self.RANK[f.suffix.lower()], str(f))):
            self.by_stem.setdefault(f.stem.lower(), f)
            for k in coverage._label_in_name(f.stem[:120]):
                self.by_key.setdefault(k, f)

    def find(self, row: dict[str, Any]) -> Path | None:
        """The file for a figure item: a name the author's script writes for it (Figure_7.pdf finds Figure_7.png too), else
        one named after its label."""
        for o in _list(row.get("outputs"), 50):
            if isinstance(o, str) and (f := self.by_stem.get(posixpath.splitext(posixpath.basename(clean(o, 200)))[0].lower())):
                return f
        return self.by_key.get(row_key(row))


class Pictures:
    """The pictures of one deck: what they may weigh in all, and when to stop rendering."""

    def __init__(self, deadline: float | None = None, total: int | None = None):
        self.left = IMAGE_TOTAL_BYTES if total is None else total
        self.deadline = deadline

    def late(self) -> bool:
        return self.deadline is not None and time.monotonic() > self.deadline

    def take(self, img: Img | None) -> bool:
        if img is None:
            return False
        if len(img.data) > self.left:
            return False
        self.left -= len(img.data)
        return True


# A PDF an author's script wrote is hostile input: a few hundred KB of flate can inflate to gigabytes and pymupdf holds the
# GIL while it works, so the whole app would stand still. The first page is rendered by a child process (the bytes go in on
# stdin, a PNG comes back through a temporary file), killed at RENDER_SECONDS; no picture is the answer to a timeout.
_RENDER_CHILD = """
import sys, pymupdf
out, max_px, dpi = sys.argv[1], float(sys.argv[2]), float(sys.argv[3])
doc = pymupdf.open(stream=sys.stdin.buffer.read(), filetype="pdf")
if doc.page_count == 0:
    sys.exit(3)
page = doc[0]
scale = min(dpi / 72.0, max_px / max(page.rect.width, page.rect.height, 1.0))
page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False).save(out)
"""


def render_pdf_figure(path: Path) -> Img | None:
    data = _read_file(path, PDF_FIGURE_BYTES)
    if data is None:
        return None
    try:
        python = W.venv_python()
    except W.WorkspaceError:
        return None
    with tempfile.TemporaryDirectory(prefix="rl-deck-") as tmp:
        out = Path(tmp) / "page.png"
        try:
            done = subprocess.run([str(python), "-I", "-c", _RENDER_CHILD, str(out), str(IMAGE_MAX_PX), str(CROP_DPI)], input=data,
                                  capture_output=True, timeout=RENDER_SECONDS, env={"PATH": "/usr/bin:/bin"})
        except (subprocess.TimeoutExpired, OSError):
            return None
        png = _read_file(out) if done.returncode == 0 else None
    return _png_bounded(png) if png else None


def figure_image(path: Path) -> Img | None:
    if path.suffix.lower() == ".pdf":
        return render_pdf_figure(path)
    data = _read_file(path)
    return _png_bounded(data) if data is not None else None


# ---------------------------------------------------------------- the slides

def _cn(row: dict[str, Any]) -> str:
    num = clean(row.get("label"), 24).rsplit(" ", 1)[-1]
    return f"{'表' if row.get('kind') == 'table' else '图'} {num}"


def _scripts_of(row: dict[str, Any]) -> list[str]:
    names = []
    for s in _list(row.get("scripts"), 10):
        n = clean(posixpath.basename(clean(_dict(s).get("path"), 160)), 60)
        if n and n not in names:
            names.append(n)
    return names


def _waiting_for(ctx: Ctx, row: dict[str, Any]) -> list[str]:
    mine = {n.lower() for n in _scripts_of(row)}
    out: dict[str, None] = {}
    for script, files in ctx.waiting.items():
        if posixpath.basename(script).lower() in mine:
            for f in _list(files, MAX_WAITING):
                out.setdefault(clean(f, 60), None)
    return list(out)[:MAX_WAITING]


def _verdict_text(counts: dict[str, int]) -> str:
    return " · ".join(f"{VERDICT_NAMES[v]} {counts[v]}" for v in ("exact", "within", "direction", "mismatch", "missing") if counts.get(v))


def result_spec(ctx: Ctx, row: dict[str, Any], crop: Crop | None, crop_note: str, figures: Figures, part: str,
                pictures: Pictures | None = None) -> Spec:
    label = clean(row.get("label"), 24)
    kind = "table" if row.get("kind") == "table" else "figure"
    where = f"{'正文' if row.get('source') == 'main' else '附录'} {_cn(row)}"
    page = clean(row.get("page"), 12)
    scripts = _scripts_of(row)
    spec = Spec("result", title=clean(f"{_cn(row)}：{clean(row.get('caption'), 240)}".rstrip("：").strip(), TITLE_CHARS), crumb=part,
                owner=clean(row.get("owner"), SHORT_CHARS))
    if crop is not None:
        spec.left = crop.img
        spec.left_note = "" if crop.mode == "caption" else f"没找到这张{'表' if kind == 'table' else '图'}的标题位置，放的是整页（论文第 {page} 页）"
    else:
        spec.left_note = crop_note or "论文里没有找到这一页的原图"
    counts: dict[str, int] = {}
    if ctx.run is not None and kind == "table":
        spec.matrix = matrix_of(ctx.tables, label, ctx.verdicts.get(_nt(label)))
        if spec.matrix:
            spec.right = "table"
            counts = ctx.counts.get(_nt(label), {})
    elif ctx.run is not None:
        pictures = pictures or Pictures()
        path = figures.find(row)
        if path is not None and pictures.late():
            spec.right, spec.right_note, spec.right_detail = "note", "出稿时间到了", "这一页没有放复现图"
        elif path is not None:
            img = figure_image(path)
            if img is not None and pictures.take(img):
                spec.right, spec.right_img = "image", img
            elif img is not None:
                spec.right, spec.right_note = "note", "图过大，未放入"
    if not spec.right:
        waiting = _waiting_for(ctx, row)
        spec.right = "note"
        if waiting:
            spec.right_note = "等数据"
            spec.right_detail = "还缺这些文件：" + "、".join(clean(w, 40) for w in waiting[:MAX_FILE_NAMES]) + (
                f" 等 {len(waiting)} 个" if len(waiting) > MAX_FILE_NAMES else "")
        else:
            spec.right_note = "尚未复现"
            if ctx.run is None:
                spec.right_detail = "还没有可用的复现运行"
            elif not scripts:
                spec.right_detail = "清单里没有对上这张的脚本"
            else:
                spec.right_detail = f"脚本 {scripts[0]} 还没有产出这张{'表' if kind == 'table' else '图'}"
    foot = f"来源：{where}" + (f"（论文第 {page} 页）" if page else "") + "；" + ("、".join(scripts[:MAX_FILE_NAMES]) if scripts else "（未对上脚本）")
    if counts:
        foot += "　判定：" + _verdict_text(counts) + ("" if ctx.reviewed else "（目标未审阅）")
    if spec.matrix and spec.matrix.notes:
        foot += "　" + spec.matrix.notes[0]
    spec.footer = foot
    return spec


def _facts(pairs: list[tuple[str, str]]) -> list[Bullet]:
    return [Bullet(f"{k}：{v}" if v else f"{k}：{MISSING}", not v) for k, v in pairs]


def _cited(summary: dict[str, Any], key: str, pages: list[str]) -> str:
    v = summary.get(key)
    if isinstance(v, dict):
        for p in _list(v.get("pages"), 4):
            p = clean(p, 8)
            if p and p not in pages:
                pages.append(p)
        return clean(v.get("text"), BULLET_CHARS)
    return clean(v, BULLET_CHARS)


def _joined(v: Any, n: int = 8, each: int = 24) -> str:
    return "、".join(x for x in (clean(i, each) for i in _list(v, n)) if x)


def _text_slide(title: str, crumb: str, bullets: list[Bullet], pages: list[str], extra: str = "") -> Spec:
    bullets = bullets[:MAX_BULLETS]
    thin = len(bullets) < 3 or any(b.missing for b in bullets)
    foot = "草稿：来自文献概要和复现方案，请核对" + (f"；论文第 {'、'.join(pages[:5])} 页" if pages else "") + (f"；{extra}" if extra else "")
    return Spec("text", title=title, crumb=crumb, bullets=bullets, thin=thin, tag="待修改" if thin else "", footer=foot)


def narrative_specs(ctx: Ctx) -> tuple[list[Spec], list[Spec]]:
    """Parts 01 and 02, from the stored summary and the plan: no new AI call."""
    s, pl = ctx.summary, _dict(ctx.plan.get("plan"))
    c1, c2 = f"{PARTS[0][0]} {PARTS[0][1]}", f"{PARTS[1][0]} {PARTS[1][1]}"
    pg: list[str] = []
    rq, arg = _cited(s, "research_question", pg), _cited(s, "argument", pg)
    first = _text_slide("研究问题", c1, _facts([("研究问题", rq), ("核心论点", arg), ("案例", _joined(s.get("cases")))]), pg)
    pg = []
    mech, usage = _cited(s, "mechanism", pg), _cited(s, "concept_usage", pg)
    second = _text_slide("机制与概念", c1, _facts([("机制", mech), ("概念用法", usage), ("关键词", _joined(s.get("key_terms")))]), pg)
    pg = []
    model = _dict(pl.get("model"))
    family = clean(model.get("family"), 30)
    spec_text = clean(model.get("specification"), BULLET_CHARS)
    model_text = clean((family + "：" if family and family != "other" else "") + spec_text, BULLET_CHARS) if spec_text or family != "other" else ""
    assumptions = [clean(_dict(a).get("text"), BULLET_CHARS) for a in _list(pl.get("assumptions"), MAX_LIST) if _dict(a).get("text")]
    bullets = _facts([("方法", _cited(s, "method", pg)), ("模型", model_text)])
    bullets += [Bullet(f"识别假设：{a}") for a in assumptions[:2]] or [Bullet(f"识别假设：{MISSING}", True)]
    if len(assumptions) > 2:
        bullets.append(Bullet(f"另有 {len(assumptions) - 2} 条识别假设，见复现方案"))
    third = _text_slide("模型设定与识别假设", c1, bullets, pg)
    part1 = [first, second, third]
    pg = []
    sources = [_dict(x) for x in _list(pl.get("data_sources"), MAX_LIST)]
    period = clean(pl.get("period"), 60) or _cited(s, "period", pg)
    bullets = _facts([("分析单位", clean(pl.get("unit_of_analysis"), BULLET_CHARS)), ("时期", period),
                      ("样本规则", clean(pl.get("sample_rule"), BULLET_CHARS))])
    for d in sources[:2]:
        name, how = clean(d.get("name"), 80), clean(d.get("access"), 20)
        if name:
            bullets.append(Bullet(f"数据：{name}" + (f"（{how}）" if how else "")))
    if len(sources) > 2:
        bullets.append(Bullet(f"另有 {len(sources) - 2} 个数据来源，见复现方案"))
    if not sources:
        bullets.append(Bullet(f"数据：{MISSING}", True))
    part2 = [_text_slide("数据来源", c2, bullets, pg)]
    variables = [_dict(v) for v in _list(pl.get("variables"), MAX_LIST) if _dict(v).get("name")]
    shown = variables[:MAX_VARIABLE_SLIDES * VARIABLES_PER_SLIDE]
    chunks = [shown[i:i + VARIABLES_PER_SLIDE] for i in range(0, len(shown), VARIABLES_PER_SLIDE)] or [[]]
    for n, chunk in enumerate(chunks, 1):
        bl = []
        for v in chunk:
            role = plans.ROLE_LABELS.get(clean(v.get("role"), 20), clean(v.get("role"), 20))
            measure = clean(v.get("measure"), 120)
            bl.append(Bullet(f"{role}：{clean(v.get('name'), 60)}" + (f"（{measure}）" if measure else "")))
        if not bl:
            bl = [Bullet(f"变量：{MISSING}", True)]
        extra = f"另有 {len(variables) - len(shown)} 个变量未列出" if n == len(chunks) and len(variables) > len(shown) else ""
        part2.append(_text_slide("变量与测量" + (f"（{n}/{len(chunks)}）" if len(chunks) > 1 else ""), c2, bl, [], extra))
    return part1, part2


def conclusion_spec(ctx: Ctx, picked: list[Spec]) -> Spec:
    s = ctx.summary
    pg: list[str] = []
    bullets = _facts([("论文的核心论点", _cited(s, "argument", pg)), ("机制", _cited(s, "mechanism", pg))])
    done = sum(1 for x in picked if x.right in ("table", "image"))
    waiting = sum(1 for x in picked if x.right_note == "等数据")
    bullets.append(Bullet(f"复现情况：勾选的 {len(picked)} 项里，已复现 {done} 项，等数据 {waiting} 项，尚未复现 {len(picked) - done - waiting} 项"
                          if picked else "复现情况：还没有勾选要复现的表和图"))
    counts = _dict(_dict(ctx.comparison).get("summary")).get("counts")
    if isinstance(counts, dict) and any(isinstance(v, int) and v for v in counts.values()):
        total = sum(v for v in counts.values() if isinstance(v, int))
        bullets.append(Bullet(f"对照论文报告值：共 {total} 项，" + "，".join(
            f"{VERDICT_NAMES[k]} {counts[k]}" for k in ("exact", "within", "direction", "mismatch", "missing") if counts.get(k))))
    return _text_slide("结论", f"{PARTS[4][0]} {PARTS[4][1]}", bullets, pg)


def blank_spec(title: str, prompts: list[str]) -> Spec:
    return Spec("blank", title=title, crumb=f"{PARTS[4][0]} {PARTS[4][1]}", tag="待填写",
                bullets=[Bullet(p, True) for p in prompts], footer="空白模板，由汇报人填写")


def authors_of(raw: Any) -> str:
    try:
        data = json.loads(raw) if isinstance(raw, str) and len(raw) < 100_000 else []
    except ValueError:
        data = []
    names = []
    for a in _list(data, 12):
        if isinstance(a, dict):
            name = clean(" ".join(x for x in (clean(a.get("given"), 40), clean(a.get("family") or a.get("last"), 40)) if x) or a.get("name"), 60)
        else:
            name = clean(a, 60)
            if "," in name:
                last, _, first = name.partition(",")
                name = f"{first.strip()} {last.strip()}".strip()
        if name:
            names.append(name)
    return ", ".join(names[:8]) + (" 等" if len(names) > 8 else "")


def presenters_of(rows: list[dict[str, Any]]) -> str:
    owners: list[str] = []
    for ticked_only in (True, False):
        for r in rows:
            o = clean(r.get("owner"), SHORT_CHARS)
            if o and (r.get("ticked") or not ticked_only) and o not in owners:
                owners.append(o)
        if owners:
            break
    return "、".join(owners[:8])


def cover_spec(ctx: Ctx) -> Spec:
    p = ctx.paper
    venue = " · ".join(x for x in (clean(p.get("venue"), 80), clean(p.get("year"), 6)) if x)
    return Spec("cover", title=clean(p.get("title") or "（未关联文献）", 160), cover={
        "authors": authors_of(p.get("authors_json")), "venue": venue, "presenters": presenters_of(ctx.rows),
        "date": ctx.seminar_date, "tag": "DCL sharing"})


def build(ctx: Ctx, pdf: PaperPdf | None = None, *, deadline: float | None = None) -> tuple[list[Spec], list[str]]:
    """Every slide of the deck, in order, at most MAX_SLIDES."""
    notes = list(ctx.notes)
    part1, part2 = narrative_specs(ctx)
    main = [r for r in ctx.rows if r.get("ticked") and r.get("group") == "main"]
    appendix = [r for r in ctx.rows if r.get("ticked") and r.get("group") == "robust"] + \
               [r for r in ctx.rows if r.get("ticked") and r.get("group") == "other"]
    fixed = 1 + len(PARTS) + len(part1) + len(part2) + 3 + (0 if main else 1) + (0 if appendix else 1)
    room = MAX_SLIDES - fixed
    if room < 0:
        raise DeckError("变量太多，排不进展示稿。")
    budget = max(0, room - (1 if len(main) + len(appendix) > room else 0))
    figures = Figures(ctx.plan_id, str(ctx.run["id"])) if ctx.run is not None else Figures()
    crumbs = {"main": f"{PARTS[2][0]} {PARTS[2][1]}", "appendix": f"{PARTS[3][0]} {PARTS[3][1]}"}
    late = False
    pictures = Pictures(deadline)

    def make_all(rows: list[dict[str, Any]], part: str) -> list[Spec]:
        nonlocal late
        out = []
        for row in rows:
            crop, why = None, ""
            if pdf is not None and pdf.files:
                if pictures.late():
                    late, why = True, "出稿时间到了，这一页没有截原图"
                else:
                    try:
                        crop = pdf.crop(row)
                    except (RuntimeError, ValueError, OSError) as e:
                        log.info("crop failed for %s: %s", row.get("label"), e)
                        why = "原图截取出错"
                if crop is not None and not pictures.take(crop.img):
                    crop, why = None, "图过大，未放入"
            out.append(result_spec(ctx, row, crop, why, figures, part, pictures))
        return out

    take_main = main[:budget]
    take_app = appendix[:max(0, budget - len(take_main))]
    skipped = [r for r in main[len(take_main):]] + [r for r in appendix[len(take_app):]]
    results3 = make_all(take_main, crumbs["main"])
    results4 = make_all(take_app, crumbs["appendix"])
    if late or any(x.right_note == "出稿时间到了" for x in results3 + results4):
        notes.append("出稿时间到了，后面的页没有截论文原图或放复现图。")
    specs = [cover_spec(ctx)]
    for n, (num, name) in enumerate(PARTS):
        specs.append(Spec("part", title=name, number=num))
        if n == 0:
            specs += part1
        elif n == 1:
            specs += part2
        elif n == 2:
            specs += results3 or [Spec("blank", title="尚未勾选正文结果", crumb=crumbs["main"], footer="在全书复现清单里勾选后重新出稿")]
        elif n == 3:
            specs += results4 or [Spec("blank", title="尚未勾选附录稳健性检验", crumb=crumbs["appendix"], footer="在全书复现清单里勾选后重新出稿")]
            if skipped:
                names = "、".join(clean(r.get("label"), 24) for r in skipped[:12]) + (" 等" if len(skipped) > 12 else "")
                specs.append(Spec("text", title=f"另有 {len(skipped)} 项未列出", crumb=crumbs["appendix"],
                                  bullets=[Bullet(f"展示稿最多 {MAX_SLIDES} 页，没排进来的：{names}")]))
        else:
            specs += [conclusion_spec(ctx, results3 + results4),
                      blank_spec("评价", ["做得好的地方", "可以商榷的地方", "对我们研究的启发"]),
                      blank_spec("疑问", ["1.", "2.", "3."])]
    if skipped:
        notes.append(f"展示稿最多 {MAX_SLIDES} 页，另有 {len(skipped)} 项没排进来。")
    return specs[:MAX_SLIDES], notes


# ---------------------------------------------------------------- the .pptx

def _inches(v: float) -> int:
    from pptx.util import Inches
    return int(Inches(v))


def _style(run: Any, size: float, color: str = INK, bold: bool = False, italic: bool = False) -> None:
    from lxml import etree
    from pptx.dml.color import RGBColor
    from pptx.oxml.ns import qn
    from pptx.util import Pt
    f = run.font
    f.size, f.bold, f.italic = Pt(size), bold, italic
    f.color.rgb = RGBColor.from_string(color)
    f.name = FONT
    rpr = run._r.get_or_add_rPr()
    rpr.set("lang", "zh-CN")
    latin = rpr.find(qn("a:latin"))
    ea = etree.Element(qn("a:ea"))
    ea.set("typeface", FONT)
    latin.addnext(ea)


def _text(slide: Any, x: float, y: float, w: float, h: float, text: str, size: float, *, color: str = INK, bold: bool = False,
          italic: bool = False, align: str = "l", anchor: str = "t") -> Any:
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    box = slide.shapes.add_textbox(_inches(x), _inches(y), _inches(w), _inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    p = tf.paragraphs[0]
    p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
    _style(p.add_run(), size, color, bold, italic)
    p.runs[0].text = text
    return box


def _rect(slide: Any, x: float, y: float, w: float, h: float, fill: str | None = None, line: str | None = None, rounded: bool = False) -> Any:
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, _inches(x), _inches(y), _inches(w), _inches(h))
    shape.shadow.inherit = False
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = RGBColor.from_string(line)
        shape.line.width = 9525
    else:
        shape.line.fill.background()
    return shape


def _tag(slide: Any, text: str, x: float = 11.43, y: float = 0.2) -> None:
    box = _rect(slide, x, y, 1.3, 0.34, TAG_FILL, rounded=True)
    tf = box.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    _style(tf.paragraphs[0].add_run(), 12, TAG_TEXT, True)
    tf.paragraphs[0].runs[0].text = text


def _page_number(slide: Any, n: int) -> None:
    """A real slide-number field, so numbers follow a slide moved in Keynote or PowerPoint."""
    from lxml import etree
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml.ns import qn
    box = slide.shapes.add_textbox(_inches(11.93), _inches(6.98), _inches(0.8), _inches(0.3))
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    fld = etree.SubElement(p._p, qn("a:fld"))
    fld.set("id", "{" + str(uuid.uuid4()).upper() + "}")
    fld.set("type", "slidenum")
    rpr = etree.SubElement(fld, qn("a:rPr"))
    rpr.set("lang", "en-US")
    rpr.set("sz", "1100")
    fill = etree.SubElement(etree.SubElement(rpr, qn("a:solidFill")), qn("a:srgbClr"))
    fill.set("val", MUTED)
    for tag in ("a:latin", "a:ea"):
        etree.SubElement(rpr, qn(tag)).set("typeface", FONT)
    etree.SubElement(fld, qn("a:t")).text = str(n)


def _bullets(slide: Any, bullets: list[Bullet], x: float, y: float, w: float, h: float, size: float = 20) -> None:
    from lxml import etree
    from pptx.oxml.ns import qn
    from pptx.util import Pt
    box = slide.shapes.add_textbox(_inches(x), _inches(y), _inches(w), _inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(12)
        p.line_spacing = 1.15
        ppr = p._p.get_or_add_pPr()
        ppr.set("marL", "300000")
        ppr.set("indent", "-300000")
        bu = etree.SubElement(ppr, qn("a:buChar"))
        bu.set("char", "•")
        run = p.add_run()
        _style(run, size, MUTED if b.missing else INK, False, b.missing)
        run.text = b.text


def _fit(img: Img, x: float, y: float, w: float, h: float) -> tuple[float, float, float, float]:
    k = min(w / img.w, h / img.h)
    iw, ih = img.w * k, img.h * k
    return x + (w - iw) / 2, y, iw, ih


def _picture(slide: Any, img: Img, x: float, y: float, w: float, h: float) -> None:
    px, py, pw, ph = _fit(img, x, y, w, h)
    slide.shapes.add_picture(io.BytesIO(img.data), _inches(px), _inches(py), _inches(pw), _inches(ph))


def _cell_text(cell: Any, c: Cell, size: float, header: bool = False, first: bool = False) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Pt
    tf = cell.text_frame
    tf.word_wrap = True
    cell.margin_left = cell.margin_right = Pt(2)
    cell.margin_top = cell.margin_bottom = Pt(1)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, line in enumerate(c.lines or [""]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT if first else PP_ALIGN.CENTER
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = 0.95
        run = p.add_run()
        _style(run, size if i == 0 else size - 1, NAVY if header else (INK if i == 0 else MUTED), header or c.bold)
        run.text = line or "\u00a0"
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor.from_string(HEAD_FILL if header else (c.fill or "FFFFFF"))


def _matrix(slide: Any, m: Matrix, x: float, y: float, w: float, h: float) -> None:
    from lxml import etree
    from pptx.oxml.ns import qn
    cols = len(m.header)
    term_w = min(1.35, w * 0.25)
    shape = slide.shapes.add_table(len(m.rows) + 1, cols, _inches(x), _inches(y), _inches(w), _inches(ROW_IN * (len(m.rows) + 1)))
    tbl = shape.table
    tbl.first_row, tbl.horz_banding = True, False
    style = shape._element.find(".//" + qn("a:tableStyleId"))
    if style is not None:
        style.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"          # No Style, No Grid
    tbl.columns[0].width = _inches(term_w)
    for j in range(1, cols):
        tbl.columns[j].width = _inches((w - term_w) / max(1, cols - 1))
    tbl.rows[0].height = _inches(ROW_IN)
    for i, c in enumerate(m.header):
        _cell_text(tbl.cell(0, i), Cell([c]), CELL_PT, header=True, first=i == 0)
    for r, row in enumerate(m.rows, 1):
        tbl.rows[r].height = _inches(ROW_IN)
        for j, c in enumerate(row):
            _cell_text(tbl.cell(r, j), c, CELL_PT, first=j == 0)
            ln = etree.Element(qn("a:lnB"))
            ln.set("w", "6350")
            etree.SubElement(etree.SubElement(ln, qn("a:solidFill")), qn("a:srgbClr")).set("val", RULE)
            tbl.cell(r, j)._tc.get_or_add_tcPr().insert(0, ln)


def _theme_fonts(prs: Any) -> None:
    """PingFang SC on the Mac; Windows has no such font, so the theme names a per-script fallback."""
    from lxml import etree
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    try:
        part = prs.slide_master.part.part_related_by(RT.THEME)
        root = part._element if hasattr(part, "_element") else etree.fromstring(part.blob)
    except (KeyError, ValueError, etree.XMLSyntaxError):
        return
    for group in ("majorFont", "minorFont"):
        for node in root.iterfind(f".//a:{group}", ns):
            for tag in ("latin", "ea"):
                el = node.find(f"a:{tag}", ns)
                if el is not None:
                    el.set("typeface", FONT)
            for el in node.iterfind("a:font", ns):
                if el.get("script") in FALLBACK_FONTS:
                    el.set("typeface", FALLBACK_FONTS[el.get("script")])
    if not hasattr(part, "_element"):
        part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _crumb(slide: Any, text: str) -> None:
    if text:
        _text(slide, 0.6, 0.22, 8, 0.3, text, 11, color=MUTED)


def _footer(slide: Any, text: str) -> None:
    if text:
        _text(slide, 0.6, 6.98, 11.2, 0.3, text, 10.5, color=MUTED)


def _title(slide: Any, text: str, size: float = 26) -> None:
    _text(slide, 0.6, 0.55, 12.13, 0.85, text, size if len(text) <= 60 else size - 4 if len(text) <= 110 else size - 6, color=NAVY, bold=True, anchor="m")
    _rect(slide, 0.6, 1.45, 12.13, 0.014, RULE)


def _draw(prs: Any, spec: Spec, n: int) -> None:
    from pptx.enum.text import MSO_ANCHOR
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    k = spec.kind
    if k == "cover":
        c = spec.cover
        _rect(slide, 0, 0, 0.25, 7.5, NAVY)
        _text(slide, 1.0, 1.6, 11.3, 1.9, spec.title, 36 if len(spec.title) <= 70 else 30, color=NAVY, bold=True, anchor="b")
        _rect(slide, 1.0, 3.7, 2.2, 0.03, NAVY)
        _text(slide, 1.0, 3.95, 11.3, 0.5, c.get("authors", ""), 20)
        _text(slide, 1.0, 4.5, 11.3, 0.45, c.get("venue", ""), 18, color=MUTED)
        _text(slide, 1.0, 5.3, 11.3, 0.4, "汇报人：" + (c.get("presenters") or "　　　　　　"), 16)
        _text(slide, 1.0, 5.75, 11.3, 0.4, "研讨日期：" + (c.get("date") or "　　　　　　"), 16)
        _text(slide, 1.0, 6.6, 11.3, 0.4, c.get("tag", ""), 14, color=MUTED, bold=True)
        return
    if k == "part":
        _text(slide, 1.0, 2.5, 4, 1.3, spec.number, 72, color=NAVY, bold=True, anchor="b")
        _rect(slide, 1.0, 3.95, 1.4, 0.03, NAVY)
        _text(slide, 1.0, 4.15, 11, 0.9, spec.title, 34, color=INK, bold=True)
        _page_number(slide, n)
        return
    _crumb(slide, spec.crumb)
    _title(slide, spec.title)
    if spec.tag:
        _tag(slide, spec.tag)
    if k in ("text", "blank"):
        if spec.bullets:
            chars = sum(len(b.text) for b in spec.bullets)
            _bullets(slide, spec.bullets, 0.8, 1.9, 11.7, 4.9, 20 if k == "blank" else 22 if chars <= 330 else 20 if chars <= 520 else 18)
    else:
        left_x, right_x, top, w, h = 0.6, 6.78, 1.95, 5.95, 4.65
        _text(slide, left_x, 1.58, 3, 0.3, "原始（论文）", 13, color=MUTED, bold=True)
        _text(slide, right_x, 1.58, 3, 0.3, "现在（复现）", 13, color=MUTED, bold=True)
        if spec.owner:
            _text(slide, 8.7, 0.22, 4.03, 0.3, f"汇报人：{spec.owner}", 11, color=MUTED, align="r")
        if spec.matrix and spec.matrix.legend:
            _text(slide, right_x + 2.2, 1.6, w - 2.2, 0.3, spec.matrix.legend, 10, color=MUTED, align="r")
        _rect(slide, left_x, top, w, h, PANEL if not spec.left else None)
        if spec.left:
            _picture(slide, spec.left, left_x, top, w, h)
            if spec.left_note:
                _text(slide, left_x, top + h + 0.02, w, 0.22, spec.left_note, 9, color=MUTED)
        else:
            _text(slide, left_x + 0.2, top, w - 0.4, h, spec.left_note, 14, color=MUTED, align="c", anchor="m")
        if spec.right == "table" and spec.matrix:
            _matrix(slide, spec.matrix, right_x, top, w, h)
        elif spec.right == "image" and spec.right_img:
            _picture(slide, spec.right_img, right_x, top, w, h)
        else:
            _rect(slide, right_x, top, w, h, PANEL)
            _text(slide, right_x + 0.2, top + h / 2 - 0.55, w - 0.4, 0.7, spec.right_note, 28, color=MUTED, bold=True, align="c", anchor="b")
            _text(slide, right_x + 0.2, top + h / 2 + 0.25, w - 0.4, 0.8, spec.right_detail, 13, color=MUTED, align="c")
    _footer(slide, spec.footer)
    _page_number(slide, n)


def write_pptx(specs: list[Spec], out: Any, title: str = "") -> None:
    """The deck as a .pptx (16:9, blank layout, every slide drawn by hand) into the path or file object `out`."""
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    _theme_fonts(prs)
    prs.core_properties.title = clean(title, 200)
    prs.core_properties.author = "Research Lab"
    for n, spec in enumerate(specs, 1):
        _draw(prs, spec, n)
    prs.save(out)


# ---------------------------------------------------------------- the file

def exports_dir(plan_id: str) -> Path:
    return W.workspace_dir(plan_id) / EXPORTS


def deck_path(plan_id: str) -> Path:
    try:
        return exports_dir(plan_id) / f"{plan_id}-deck.pptx"
    except W.WorkspaceError as e:
        raise DeckError(str(e)) from None


def info(plan_id: str) -> dict[str, Any]:
    """The deck made last: {exists, name, size, when}. Anything but one plain file (a link, a second hard link) is not it."""
    try:
        path = deck_path(plan_id)
        W._checked(W.workspace_dir(plan_id), path)
        st = os.lstat(path)
    except (DeckError, W.WorkspaceError, OSError):
        return {"exists": False, "name": "", "size": 0, "when": ""}
    if not stat.S_ISREG(st.st_mode) or st.st_nlink > 1:
        return {"exists": False, "name": "", "size": 0, "when": ""}
    return {"exists": True, "name": path.name, "size": st.st_size, "when": time.strftime("%Y-%m-%d %H:%M", time.localtime(st.st_mtime))}


def download_path(plan_id: str) -> Path | None:
    """The one file the download serves, or None."""
    return deck_path(plan_id) if info(plan_id)["exists"] else None


def _save(plan_id: str, specs: list[Spec], path: Path, title: str) -> None:
    ws = W.workspace_dir(plan_id)
    try:
        W._checked(ws, path)
        path.parent.mkdir(parents=True, exist_ok=True)
        W._checked(ws, path.parent / "x")
    except W.WorkspaceError as e:
        raise DeckError(str(e)) from None
    tmp = path.with_name(f".{path.name}.{os.getpid()}.{uuid.uuid4().hex[:8]}.tmp")
    fd = os.open(tmp, os.O_WRONLY | os.O_CREAT | os.O_EXCL | os.O_NOFOLLOW, 0o644)
    try:
        with os.fdopen(fd, "wb") as fh:
            write_pptx(specs, fh, title)
        os.replace(tmp, path)
    except BaseException:
        try:
            os.unlink(tmp)
        except OSError:
            pass
        raise


def make(conn: sqlite3.Connection, plan_id: str, *, seminar_date: str = "") -> DeckResult:
    """Build the deck of a plan (deck v2, card 0196: report_content assembles it, slides draws it) and save it as
    <workspace>/exports/<plan>-deck.pptx, replacing the last one, with its pictures beside it for the preview page."""
    from . import report_content, slides
    path = deck_path(plan_id)
    built = report_content.build(conn, plan_id, deadline=time.monotonic() + DECK_SECONDS, seminar_date=seminar_date)
    # the pptx carries Leo's edits; slides.json keeps the slides as made, so the edits stay keyed to them
    import copy
    edited = report_content.apply_edits(copy.deepcopy(built.deck), report_content.load_edits(plan_id))
    _save_v2(plan_id, edited, path)
    try:
        report_content.save_assets(plan_id, built.deck)
    except (OSError, W.WorkspaceError) as e:
        built.notes.append(f"预览图没有存下：{clean(str(e), 80)}")
    return DeckResult(path, len(built.deck.slides), built.notes, built.replicated, built.pending)


def _save_v2(plan_id: str, deck_v2: Any, path: Path) -> None:
    from . import slides
    ws = W.workspace_dir(plan_id)
    try:
        W._checked(ws, path)
        path.parent.mkdir(parents=True, exist_ok=True)
        W._checked(ws, path.parent / "x")
    except W.WorkspaceError as e:
        raise DeckError(str(e)) from None
    tmp = path.with_name(f".{path.name}.{os.getpid()}.{uuid.uuid4().hex[:8]}.tmp")
    fd = os.open(tmp, os.O_WRONLY | os.O_CREAT | os.O_EXCL | os.O_NOFOLLOW, 0o644)
    try:
        with os.fdopen(fd, "wb") as fh:
            slides.write_pptx(deck_v2, fh)
        os.replace(tmp, path)
    except BaseException:
        try:
            os.unlink(tmp)
        except OSError:
            pass
        raise

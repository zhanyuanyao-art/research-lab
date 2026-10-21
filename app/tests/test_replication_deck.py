"""Card 0190 (docs/design/REPLICATION-v4-DCL.md §3D): the 展示稿. A .pptx in the DCL lab's parts, built by rules from the
stored summary, the plan, the 全书复现清单 and the latest valid run: left a crop of the paper, right our table or figure,
verdicts from compare, 「尚未复现」 / 「等数据」 where there is nothing, 评价 and 疑问 left empty."""
from __future__ import annotations

import io
import json
import os
import threading
import time
import zlib
from pathlib import Path
from typing import Any

import pymupdf
import pytest

from research_lab.db.connection import next_sequence_id, now_iso
from research_lab.replication import coverage as cov
from research_lab.replication import deck, plans
from research_lab.replication import workspace as W

from .test_data import data_client, ingest_quiet  # noqa: F401 - data_client is a fixture

LONG = "Advisers shape the decisions that national security meetings reach, and the chair listens to them closely. " * 3


def layout_pdf(path: Path, pages: list[list[tuple]], rotate: int = 0) -> Path:
    """Pages of ("t", x, y, text[, size]) lines and ("r", x0, y0, x1, y1) boxes, so a test controls every gap."""
    doc = pymupdf.open()
    for items in pages:
        page = doc.new_page(width=612, height=792)
        for it in items:
            if it[0] == "t":
                page.insert_text((it[1], it[2]), it[3], fontsize=it[4] if len(it) > 4 else 10)
            else:
                page.draw_rect(pymupdf.Rect(*it[1:5]), color=(0, 0, 0), fill=(0.8, 0.85, 0.95))
        if rotate:
            page.set_rotation(rotate)
    doc.save(path)
    doc.close()
    return path


def rows_text(y0: float, n: int) -> list[tuple]:
    return [("t", 72, y0 + 13 * i, f"Row {i}   0.{i}31   0.1{i}8   0.150   0.120") for i in range(n)]


MAIN_PAGES = [
    [("t", 72, 72, "Advisers and Aggregation in Foreign Policy Decision Making", 14), ("t", 72, 100, "Abstract"),
     ("t", 72, 120, "We study who decides in national security meetings.")],
    [("t", 72, 80, "4. Results"), ("t", 72, 110, "As Table 3 shows, hawkish participants matter more than others."),
     ("t", 72, 200, "TABLE 3. Effect of participant hawkishness on foreign policy decisions", 11), *rows_text(222, 6),
     ("t", 72, 305, "Notes: standard errors in parentheses."),
     # a gap of 60pt, then running text that the crop must not take
     *[("t", 72, 365 + 12 * i, LONG[:90]) for i in range(6)]],
    [("t", 72, 100, "Some text above the figure."), ("r", 100, 140, 500, 330),
     ("t", 72, 350, "FIGURE 7. Summary of the results of three models of trait aggregation", 11),
     ("t", 72, 420, "TABLE 5. Effect of principals' hawkishness on disputes, using monthly data", 11), *rows_text(442, 3),
     ("t", 72, 700, "References")],
]
APPENDIX_PAGES = [
    [("t", 72, 80, "Online Appendix"), ("t", 72, 110, "5.2 Removing the Soviet Union"),
     ("t", 72, 140, "Table A21: Effect of mean hawkishness, Removing Decisions Involving the USSR", 11), *rows_text(162, 4)],
]

SUMMARY = {"research_question": {"text": "顾问是否系统性地影响外交政策？", "pages": ["1"]},
           "argument": {"text": "顾问的倾向会进入决策。", "pages": ["2", "5"]}, "mechanism": {"text": None, "pages": []},
           "cases": ["美国国家安全委员会"], "key_terms": ["顾问", "聚合"], "period": {"text": None, "pages": []},
           "method": {"text": "生存分析与回归", "pages": ["4"]}, "concept_usage": {"text": "", "pages": []}}


def results_json(table3: bool = True) -> dict[str, Any]:
    def row(term, est, se, p, n=None):
        return {"term": term, "estimate": est, "se": se, "p": p, "ci": None, "n": n}
    tables = []
    if table3:
        for m, shift in (("(1)", 0.0), ("(2)", 0.05)):
            tables.append({"label": "Table 3", "model": m, "estimator": "OLS", "notes": [], "needs_r": False, "rows": [
                row("Mean hawkishness", 0.231 + shift, 0.054, 0.001), row("President hawkishness", 0.1 + shift, 0.04, 0.03),
                row("Observations", 1204, None, None)]})
    tables.append({"label": "Table 5", "model": "(1)", "estimator": "NB", "notes": [], "needs_r": True, "rows": []})
    return {"tables": tables}


def add_run(conn, plan_id: str, results: dict[str, Any], figures: dict[str, bytes] | None = None, valid: bool = True) -> str:
    run_id = next_sequence_id(conn, "analysis_runs", "RUN")
    conn.execute("INSERT INTO analysis_runs(id, dataset_id, plan_id, method, params_json, result_json, code, status, workspace_json, "
                 "tokens, created_at) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                 (run_id, None, plan_id, "workspace", "{}", json.dumps({"valid": valid, "results": results, "mapping": None, "errors": []}),
                  None, "done", json.dumps({"state": "done", "mode": "rrun", "code": [], "errors": []}), 0, now_iso()))
    out = W.run_dir(plan_id, run_id) / "outputs" / "figures"
    out.mkdir(parents=True)
    for name, data in (figures or {}).items():
        (out / name).write_bytes(data)
    return run_id


def png_bytes(w: int = 120, h: int = 80) -> bytes:
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, w, h), False)
    pix.set_rect(pix.irect, (40, 120, 200))
    return pix.tobytes("png")


@pytest.fixture()
def fixture_plan(conn, tmp_path) -> dict[str, Any]:
    from research_lab.intake import pipeline
    paper_id = ingest_quiet(conn, layout_pdf(tmp_path / "jost.pdf", MAIN_PAGES))
    pipeline.attach_file(conn, layout_pdf(tmp_path / "appendix.pdf", APPENDIX_PAGES), paper_id, ai=False)
    conn.execute("UPDATE papers SET title = ?, authors_json = ?, venue = ?, year = ?, ai_summary_json = ? WHERE id = ?",
                 ("Advisers and Aggregation in Foreign Policy Decision Making", json.dumps(["Jost, Tyler", "Kertzer, Joshua D."]),
                  "International Organization", 2024, json.dumps(SUMMARY, ensure_ascii=False), paper_id))
    plan_id = plans.create_blank(conn, paper_id)
    plans.save_variables(conn, plan_id, [{"role": "DV", "name": "外交政策决策", "measure": "会议结果编码"},
                                         {"role": "IV", "name": "顾问鹰派程度", "measure": "机器学习预测"}])
    plans.save_sources(conn, plan_id, [{"name": "国安会会议档案", "access": "公开"}])
    plans.save_fields(conn, plan_id, {"unit_of_analysis": "会议", "model_family": "OLS", "model_specification": "决策 ~ 鹰派程度"})
    plans.save_targets(conn, plan_id, [
        {"table": "Table 3", "model": "(1)", "term": "Mean hawkishness", "value": "0.231", "page": "2"},
        {"table": "Table 3", "model": "(1)", "term": "President hawkishness", "value": "-0.900", "page": "2"},
        {"table": "Table 3", "model": "(2)", "term": "Mean hawkishness", "value": "0.350", "page": "2"}])
    plans.review_targets(conn, plan_id)
    c = cov.rescan(conn, plan_id)
    assert c["total"] >= 4
    return {"id": plan_id, "paper": paper_id}


def read(path: Path):
    from pptx import Presentation
    return Presentation(path if hasattr(path, "read") else str(path))


def texts(slide) -> list[str]:
    out = []
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            out.append(sh.text_frame.text)
    return out


def titles(prs) -> list[str]:
    """A slide's title: the text box at the title position; for the cover and a part's divider, the biggest text that is
    not a part number."""
    from pptx.util import Inches
    out = []
    for s in prs.slides:
        top = [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip() and abs(sh.top - int(Inches(0.55))) < 100]
        if top:
            out.append(top[0].text_frame.text)
            continue
        best = ("", 0)
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                size = max((r.font.size or 0 for p in sh.text_frame.paragraphs for r in p.runs), default=0)
                if best[1] < size < 70 * 12700:
                    best = (sh.text_frame.text, size)
        out.append(best[0])
    return out


def result_slides(prs):
    return [s for s in prs.slides if any("原始" in t for t in texts(s))]


def make(conn, plan_id: str, **kw):
    return deck.make(conn, plan_id, **kw)


# ---------------------------------------------------------------- structure











# ---------------------------------------------------------------- results











def test_a_pdf_figure_is_rendered_to_a_picture(tmp_path):
    doc = pymupdf.open()
    doc.new_page(width=300, height=200).insert_text((20, 100), "plot", fontsize=20)
    f = tmp_path / "coefPlotMain.pdf"
    doc.save(f)
    img = deck.figure_image(f)
    assert img is not None and img.data[:8] == b"\x89PNG\r\n\x1a\n" and img.w > 100
    big = tmp_path / "big.png"
    big.write_bytes(png_bytes(4000, 100))
    scaled = deck.figure_image(big)
    assert scaled is not None and scaled.w <= deck.IMAGE_MAX_PX
    assert deck.figure_image(tmp_path / "missing.png") is None
    (tmp_path / "bad.png").write_bytes(b"not an image")
    assert deck.figure_image(tmp_path / "bad.png") is None


def bomb_pdf(path: Path, mb: int = 200) -> Path:
    """A PDF of about a megabyte whose one page inflates to `mb` MB of drawing operators: rendering it takes seconds, with the
    GIL held, in whichever process does it."""
    comp = zlib.compressobj(1)
    chunk = b"q Q " * (256 * 1024)
    body = b"".join(comp.compress(chunk) for _ in range(mb)) + comp.flush()
    head = (b"%PDF-1.4\n1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj\n"
            b"3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 200 200]/Contents 4 0 R>>endobj\n")
    path.write_bytes(head + b"4 0 obj<</Length " + str(len(body)).encode() + b"/Filter/FlateDecode>>stream\n" + body
                     + b"\nendstream endobj\ntrailer<</Root 1 0 R/Size 5>>\n%%EOF\n")
    return path


def test_a_pdf_figure_bomb_is_killed_and_the_app_never_stands_still(tmp_path, monkeypatch):
    monkeypatch.setattr(deck, "RENDER_SECONDS", 1.0)
    f = bomb_pdf(tmp_path / "Figure_7.pdf")
    assert f.stat().st_size < 3 * 1024 * 1024
    ticks: list[float] = []
    stop = threading.Event()

    def beat():
        while not stop.is_set():
            ticks.append(time.perf_counter())
            time.sleep(0.02)
    thread = threading.Thread(target=beat)
    thread.start()
    t = time.perf_counter()
    img = deck.figure_image(f)
    took = time.perf_counter() - t
    stop.set()
    thread.join()
    assert img is None and took < 3                                      # rendered in the app it would be ~2.5 s with nothing else running
    assert max(b - a for a, b in zip(ticks, ticks[1:])) < 0.5            # another thread of the app ran the whole time (the GIL was free)
    ok = tmp_path / "ok.pdf"
    doc = pymupdf.open()
    doc.new_page(width=200, height=100).insert_text((20, 50), "plot")
    doc.save(ok)
    assert deck.figure_image(ok) is not None                              # the same child renders an honest figure


def test_a_pdf_figure_over_the_size_cap_is_not_read_or_rendered(tmp_path):
    f = tmp_path / "big.pdf"
    f.write_bytes(b"%PDF-1.4\n" + b"\0" * deck.PDF_FIGURE_BYTES)
    assert under(1, deck.figure_image, f) is None
    link = tmp_path / "link.pdf"
    os.symlink(f, link)
    assert deck.figure_image(link) is None                                # never through a link


def test_a_picture_of_too_many_pixels_is_refused(tmp_path):
    from PIL import Image
    f = tmp_path / "huge.png"
    Image.new("1", (5000, 4000)).save(f)                                  # 20 MP, a few KB on disk
    assert deck.IMAGE_MAX_PIXELS == 16_000_000 and under(1, deck.figure_image, f) is None


def test_a_picture_wins_over_a_pdf_of_the_same_figure(conn, fixture_plan):
    run_id = add_run(conn, fixture_plan["id"], results_json(), {"Figure_7.pdf": b"%PDF-1.4", "Figure_7.png": png_bytes(), "other.pdf": b"x"})
    figures = deck.Figures(fixture_plan["id"], run_id)
    row = {"kind": "figure", "label": "Figure 7", "outputs": ["figs/Figure_7.pdf"]}
    assert figures.find(row).name == "Figure_7.png"                       # the script wrote a PDF, the run also has the PNG
    assert figures.find({"kind": "figure", "label": "Figure 7", "outputs": []}).name == "Figure_7.png"
    assert figures.find({"kind": "figure", "label": "Figure 8", "outputs": []}) is None










def test_slides_are_capped_and_the_rest_is_counted():
    rows = [{"label": f"Table {i}", "kind": "table", "caption": f"Caption {i}", "page": "1", "source": "main", "group": "main",
             "scripts": [], "outputs": [], "ticked": True, "owner": ""} for i in range(1, 121)]
    ctx = deck.Ctx({"id": "RP-2026-9999", "paper_id": None, "plan": {}}, {}, {}, rows, True)
    specs, notes = deck.build(ctx)
    assert len(specs) == deck.MAX_SLIDES
    assert [s.title for s in specs[-3:]] == ["结论", "评价", "疑问"]
    extra = next(s for s in specs if s.title.startswith("另有 "))
    assert "未列出" in extra.title and any("没排进来" in n for n in notes)
    assert sum(1 for s in specs if s.kind == "result") == deck.MAX_SLIDES - (len(specs) - sum(1 for s in specs if s.kind == "result"))


# ---------------------------------------------------------------- the crop

def open_pdf(path: Path):
    doc = pymupdf.open(path)
    return doc, [(pymupdf.Page and (doc[i].get_label() or str(i + 1))) for i in range(doc.page_count)]


def row_of(label: str, kind: str, page: str, caption: str, source: str = "main") -> dict[str, Any]:
    return {"label": label, "kind": kind, "page": page, "caption": caption, "source": source}


def crop_size(png: bytes) -> tuple[int, int]:
    pix = pymupdf.Pixmap(png)
    return pix.width, pix.height


def test_the_crop_runs_from_the_caption_to_the_next_gap(tmp_path):
    doc, labels = open_pdf(layout_pdf(tmp_path / "p.pdf", MAIN_PAGES))
    crop = deck.crop_original(doc, labels, row_of("Table 3", "table", "2", "Effect of participant hawkishness on foreign policy decisions"))
    assert crop is not None and crop.mode == "caption" and crop.page == 1
    w, h = crop_size(crop.img.data)
    # the caption (y 200) down to the notes line (y 305), about 115pt at 150 dpi, not the 700pt page and not the paragraph after it
    assert 100 / 72 * 150 < h < 150 / 72 * 150 and w < 612 / 72 * 150


def test_a_figure_whose_caption_sits_below_it_is_cropped_upward(tmp_path):
    doc, labels = open_pdf(layout_pdf(tmp_path / "p.pdf", MAIN_PAGES))
    crop = deck.crop_original(doc, labels, row_of("Figure 7", "figure", "3", "Summary of the results of three models of trait aggregation"))
    assert crop is not None and crop.mode == "caption"
    _, h = crop_size(crop.img.data)
    assert h > (330 - 140 + 20) / 72 * 150                     # the box (140 to 330) and the caption (350)
    assert h < (360 - 130) / 72 * 150 + 10                     # and not the text above it


def test_the_crop_falls_back_to_the_whole_page_when_the_caption_is_not_found(tmp_path):
    doc, labels = open_pdf(layout_pdf(tmp_path / "p.pdf", MAIN_PAGES))
    crop = deck.crop_original(doc, labels, row_of("Table 9", "table", "1", "A caption the page does not have"))
    assert crop is not None and crop.mode == "page"
    w, h = crop_size(crop.img.data)
    assert abs(w - 612 / 72 * deck.PAGE_DPI) < 3 and abs(h - 792 / 72 * deck.PAGE_DPI) < 3
    assert deck.crop_original(doc, labels, row_of("Table 9", "table", "77", "x")) is None          # a page the file has not got


def sideways_pdf(path: Path, page_rotation: int, text_rotation: int) -> Path:
    """A table page with the page's /Rotate and the way its text is drawn chosen apart: text turned 90 (reads upward, lines
    go right) or 270 (reads downward, lines go left) in the content stream, as a landscape table is set."""
    doc = pymupdf.open()
    doc.new_page(width=612, height=792).insert_text((72, 80), "Intro")
    page = doc.new_page(width=612, height=792)
    lines = ["TABLE 3. Effect of participant hawkishness on foreign policy decisions"] + [
        f"Row {i}   0.{i}31   0.1{i}8   0.150   0.120   0.33{i}   0.4{i}" for i in range(8)]
    for i, text in enumerate(lines):
        if text_rotation == 0:
            page.insert_text((72, 200 + 13 * i), text, fontsize=10)
        elif text_rotation == 90:
            page.insert_text((200 + 13 * i, 700), text, fontsize=10, rotate=90)
        else:
            page.insert_text((500 - 13 * i, 100), text, fontsize=10, rotate=270)
    if page_rotation:
        page.set_rotation(page_rotation)
    doc.save(path)
    doc.close()
    return path


@pytest.mark.parametrize("page_rotation", [0, 90, 180, 270])
@pytest.mark.parametrize("text_rotation", [0, 90, 270])
def test_a_page_that_lies_sideways_is_cropped_as_it_reads(tmp_path, page_rotation, text_rotation):
    doc, labels = open_pdf(sideways_pdf(tmp_path / "p.pdf", page_rotation, text_rotation))
    crop = deck.crop_original(doc, labels, row_of("Table 3", "table", "2", "Effect of participant hawkishness on foreign policy decisions"))
    assert crop is not None and crop.mode == "caption"
    w, h = crop_size(crop.img.data)
    # the same wide, short table whichever way the page was turned (an upright one is 674 x 272 at 150 dpi)
    assert abs(w - 674) < 8 and abs(h - 272) < 8, (w, h)
    whole = deck.crop_original(doc, labels, row_of("Table 9", "table", "2", "Not on the page"))
    assert whole is not None and whole.mode == "page"
    ww, wh = crop_size(whole.img.data)
    assert (ww > wh) == (text_rotation != 0)                               # a landscape table's page comes out landscape, upright




# ---------------------------------------------------------------- the file and the route

def test_the_file_is_written_by_rename_inside_the_workspace_only(conn, fixture_plan):
    made = make(conn, fixture_plan["id"])
    ws = W.workspace_dir(fixture_plan["id"])
    assert made.path.parent == ws / "exports" and sorted(p.name for p in (ws / "exports").iterdir()) == sorted(["deck", made.path.name])
    again = make(conn, fixture_plan["id"])
    assert again.path == made.path and len(list((ws / "exports").iterdir())) == 2
    info = deck.info(fixture_plan["id"])
    assert info["exists"] and info["name"] == made.path.name and info["size"] == made.path.stat().st_size
    with pytest.raises(deck.DeckError):
        deck.deck_path("../etc")


def test_a_link_where_the_deck_goes_is_never_followed(conn, fixture_plan, tmp_path):
    ws = W.workspace_dir(fixture_plan["id"])
    outside = tmp_path / "outside"
    outside.mkdir()
    (ws).mkdir(parents=True, exist_ok=True)
    os.symlink(outside, ws / "exports")
    with pytest.raises(deck.DeckError):
        make(conn, fixture_plan["id"])
    assert list(outside.iterdir()) == [] and not deck.info(fixture_plan["id"])["exists"]
    os.unlink(ws / "exports")
    (ws / "exports").mkdir()
    secret = tmp_path / "secret.pptx"
    secret.write_bytes(b"secret")
    os.symlink(secret, ws / "exports" / f"{fixture_plan['id']}-deck.pptx")
    assert deck.download_path(fixture_plan["id"]) is None


def test_make_deck_and_download_over_http(data_client, conn, fixture_plan):
    plan_id = fixture_plan["id"]
    r = data_client.get(f"/data/plans/{plan_id}/detail")
    assert r.status_code == 200 and "Make deck" in r.text and f"/data/plans/{plan_id}/deck" in r.text
    assert data_client.get(f"/data/plans/{plan_id}/deck.pptx").status_code == 404
    page = data_client.get(f"/data/plans/{plan_id}/deck")
    assert page.status_code == 200 and "还没有展示稿" in page.text
    r = data_client.post(f"/data/plans/{plan_id}/deck", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"].startswith(f"/data/plans/{plan_id}/deck?msg=")
    page = data_client.get(f"/data/plans/{plan_id}/deck")
    assert page.status_code == 200 and 'data-testid="deck-preview"' in page.text and "研究问题与理论" in page.text
    assert f"/data/plans/{plan_id}/deck.pptx" in page.text and f"/data/plans/{plan_id}/deck/asset/cover.png" in page.text
    assert data_client.get(f"/data/plans/{plan_id}/deck/asset/cover.png").headers["content-type"] == "image/png"
    for bad in ("..%2Fslides.json", "slides.json", "cover.PNG", "a.png.txt", "%2E%2E%2Fx.png"):
        assert data_client.get(f"/data/plans/{plan_id}/deck/asset/{bad}").status_code == 404
    got = data_client.get(f"/data/plans/{plan_id}/deck.pptx")
    assert got.status_code == 200 and got.headers["content-type"] == deck.MIME
    assert got.headers["content-disposition"].startswith("attachment") and f"{plan_id}-deck.pptx" in got.headers["content-disposition"]
    assert got.content[:2] == b"PK" and len(read(io.BytesIO(got.content)).slides) >= 6
    assert data_client.get("/data/plans/RP-2099-0001/deck.pptx").status_code == 404
    assert data_client.post("/data/plans/RP-2099-0001/deck").status_code == 404
    assert data_client.get("/data/plans/RP-2099-0001/deck").status_code == 404


def test_a_post_from_another_site_is_refused(data_client, fixture_plan):
    plan_id = fixture_plan["id"]
    headers = {"HX-Request": "", "Origin": "http://evil.example"}
    bad = data_client.post(f"/data/plans/{plan_id}/deck", headers=headers)
    assert bad.status_code == 403 and not deck.info(plan_id)["exists"]


def test_the_download_route_has_no_way_to_name_another_file(data_client, fixture_plan):
    plan_id = fixture_plan["id"]
    assert data_client.post(f"/data/plans/{plan_id}/deck", follow_redirects=False).status_code == 303
    ws = W.workspace_dir(plan_id)
    (ws / "secret.txt").write_text("TOPSECRET-CONTENT")
    for url in (f"/data/plans/{plan_id}/deck.pptx/../../secret.txt", f"/data/plans/{plan_id}/deck.pptx?file=../secret.txt",
                f"/data/plans/{plan_id}/exports/secret.txt"):
        r = data_client.get(url)
        assert b"TOPSECRET-CONTENT" not in r.content and (r.status_code == 404 or r.headers["content-type"] == deck.MIME)


# ---------------------------------------------------------------- text from outside, bounded

def under(limit: float, fn, *args, **kw):
    t = time.perf_counter()
    out = fn(*args, **kw)
    assert time.perf_counter() - t < limit
    return out


def test_clean_drops_control_and_bidi_characters_and_cuts():
    assert deck.clean("a\x00b\x07c‮d\ud800e\n\tf  g") == "a b c d e f g"
    assert deck.clean("x" * 500, 20) == "x" * 19 + "…"
    assert deck.clean(None) == "" and deck.clean({"a": 1}) == "" and deck.clean(2024, 6) == "2024"




def test_megabyte_inputs_are_cut_before_anything_reads_them():
    huge = "表 3 " * 250_000                                         # ~1 MB of one line
    under(1, deck.clean, huge)
    under(1, deck.starts_with_label, huge, "table 3")
    under(1, deck.needles_for, "table", huge)
    summary = {k: {"text": huge, "pages": [huge] * 50_000} for k in ("research_question", "argument", "mechanism", "method", "concept_usage")}
    summary["cases"] = [huge] * 50_000
    summary["key_terms"] = [huge] * 50_000
    plan_json = {"variables": [{"role": "DV", "name": huge, "measure": huge}] * 50_000, "data_sources": [{"name": huge}] * 50_000,
                 "assumptions": [{"text": huge}] * 50_000, "model": {"family": huge, "specification": huge}, "unit_of_analysis": huge}
    ctx = deck.Ctx({"id": "RP-2026-9999", "paper_id": None, "plan": plan_json}, {"title": huge, "authors_json": json.dumps([huge] * 5_000)},
                   summary, [{"label": huge, "kind": "table", "caption": huge, "page": huge, "source": "main", "group": "main",
                              "scripts": [{"path": huge}] * 5_000, "outputs": [huge] * 5_000, "ticked": True, "owner": huge}] * 2_000, True)
    specs, _ = under(1, deck.build, ctx)
    assert len(specs) <= deck.MAX_SLIDES
    longest = max(len(b.text) for s in specs if not s.title.startswith("另有") for b in s.bullets)
    assert longest <= deck.BULLET_CHARS + 20
    assert max(len(s.title) for s in specs if s.kind != "cover") <= deck.TITLE_CHARS and max(len(s.footer) for s in specs) < 1_000


def test_a_hostile_result_table_is_capped(tmp_path):
    huge = "x" * 1_000_000
    rows = [{"term": f"t{i}{huge[:2000]}", "estimate": 1.0, "se": 0.1, "p": 0.01, "ci": None, "n": None} for i in range(200_000)]
    tables = [{"label": "Table 3", "model": f"({m}){huge[:5000]}", "rows": rows[:2000]} for m in range(100)]
    m = under(1, deck.matrix_of, tables, "Table 3", None)
    assert m is not None and len(m.header) == deck.MAX_TABLE_COLS + 1 and len(m.rows) == deck.MAX_TABLE_ROWS
    assert m.notes and "未列出" in m.notes[0]
    assert all(len(c.lines[0]) <= deck.TERM_CHARS for r in m.rows for c in r[:1])


def test_page_geometry_is_linear_in_the_blocks():
    pieces = [deck.Piece(72, 200 + i * 0.001, 540, 200.0005 + i * 0.001) for i in range(200_000)]
    box = under(1, deck.crop_box, (72, 190, 540, 199), pieces, (612, 792), "figure")
    assert box[3] >= 199
    prose = [deck.Piece(72, 300.0, 540, 310.0, True)]
    far = under(1, deck.crop_box, (72, 190, 540, 199), [deck.Piece(72, 400, 540, 410), *prose], (612, 792), "table")
    assert far[3] == 199                                           # nothing within a gap of the caption: the crop is the caption
    wall = under(1, deck.crop_box, (72, 190, 540, 199), [deck.Piece(72, 205, 540, 215, True), deck.Piece(72, 216, 540, 230)], (612, 792), "table")
    assert wall[3] == 199                                          # a paragraph or the next caption right below is a wall


def test_a_page_with_a_megabyte_line_is_cropped_in_time(tmp_path):
    doc = pymupdf.open()
    page = doc.new_page(width=612, height=792)
    page.insert_text((72, 100), "TABLE 3. Effect of x on y", fontsize=10)
    page.insert_text((72, 130), "w" * 200_000, fontsize=1)           # one huge line the layout must cut
    page.insert_text((72, 140), "0.1 " * 40, fontsize=5)
    f = tmp_path / "hostile.pdf"
    doc.save(f)
    doc.close()
    d, labels = open_pdf(f)
    crop = under(1, deck.crop_original, d, labels, row_of("Table 3", "table", "1", "Effect of x on y"))
    assert crop is not None


def test_a_page_busy_with_thousands_of_runs_is_cropped_in_time(tmp_path):
    doc = pymupdf.open()
    page = doc.new_page(width=612, height=792)
    page.insert_text((72, 100), "TABLE 3. Effect of x on y", fontsize=10)
    for i in range(1_500):
        page.insert_text((72 + (i % 7) * 60, 120 + (i % 300) * 1.5), f"w{i}", fontsize=1)
    page.insert_text((10, 10), "x\n" * 30_000, fontsize=0.3)
    f = tmp_path / "busy.pdf"
    doc.save(f)
    doc.close()
    d, labels = open_pdf(f)
    crop = under(1, deck.crop_original, d, labels, row_of("Table 3", "table", "1", "Effect of x on y"))
    assert crop is not None and crop.mode in ("caption", "page")

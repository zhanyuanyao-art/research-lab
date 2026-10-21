"""Card 0196: deck v2's two renderers — the HTML preview and the .pptx — draw the same slides on the same grid."""
from __future__ import annotations

import io
import os

import pymupdf
import pytest

from research_lab.replication import method_notes, slides as S
from research_lab.replication.deck import Img


def png(w=200, h=120) -> Img:
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, w, h), False)
    pix.set_rect(pix.irect, (40, 120, 200))
    return Img(pix.tobytes("png"), w, h)


def every_kind() -> S.Deck:
    card = S.Card(head="大家的平均", body="决策反映平均倾向", tag="涌现模型", note="平均鹰派程度")
    slides = [
        S.Slide("cover", title="Advisers <script>alert(1)</script>", subtitle="顾问与外交决策", kicker="论文复现",
                meta=["Tyler Jost", "IO · 2024"], thumb="cover.png"),
        S.Slide("section", title="研究问题与理论", no="第一部分", sub="谁的倾向决定决策", topics=["问题", "论点"]),
        S.Slide("cards", crumb="第一部分 · 研究问题与理论", title="一群人开会", lede="与会者的倾向", cards=[card] * 3),
        S.Slide("flow", crumb="c", title="怎么测", cards=[card] * 3, src="复现包直接给了分数"),
        S.Slide("method", crumb="c", title="泊松回归", cards=[S.Card(head="怎么读", points=["系数", "发生率比"]), card],
                formula="log μ_i = X_i′β", code='fit = glm(y ~ x, family = "poisson")', code_head="main.R · 第 1 行"),
        S.Slide("table", crumb="c", title="复现概览", columns=["论文", "内容", "脚本", "结果"],
                rows=[["正文表 3", "x", "a.R", "完全一致"], ["正文表 4", "y", "a.R", "缺依赖"]], chips=["ok", "no"]),
        S.Slide("compare", crumb="c", title="表 3", image="orig-0.png", image_label="论文原表 · 第 26 页",
                columns=["变量", "模型", "论文", "复现"], rows=[["Mean", "(1)", "5.101", "5.101"]], hits=[True],
                verdict="完全一致"),
        S.Slide("compare", crumb="c", title="表 4", image="orig-1.png", verdict="缺依赖", verdict_ok=False,
                verdict_note="缺 glmer.nb"),
        S.Slide("compare", crumb="c", title="图 7", image="orig-0.png", image2="ours-0.png"),
        S.Slide("chart", crumb="c", title="表 3 · 对照图", image="chart-0.png"),
        S.Slide("questions", crumb="c", title="结论与疑问", cards=[S.Card(head="结论", points=["a", "b"])],
                points=["发言多就等于影响大吗", "外推可靠吗"]),
        S.Slide("points", crumb="c", title="结论", points=["顾问模型最能解释"]),
        S.Slide("cards", crumb="c", title="空的卡片页"),
    ]
    assets = {n: png() for n in ("cover.png", "orig-0.png", "orig-1.png", "ours-0.png", "chart-0.png")}
    return S.Deck("Advisers", slides, assets)


def test_the_pptx_has_one_slide_per_slide_and_carries_the_text(tmp_path):
    from pptx import Presentation
    d = every_kind()
    buf = io.BytesIO()
    S.write_pptx(d, buf)
    prs = Presentation(io.BytesIO(buf.getvalue()))
    assert len(prs.slides) == len(d.slides)
    all_text = "\n".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    for want in ("顾问与外交决策", "研究问题与理论", "决策反映平均倾向", "log μ_i = X_i′β", "缺 glmer.nb", "发言多就等于影响大吗"):
        assert want in all_text
    pics = [sum(1 for sh in s.shapes if sh.shape_type == 13) for s in prs.slides]
    assert pics[0] == 1 and pics[6] == 1 and pics[8] == 2 and pics[9] == 1


def test_the_html_escapes_outside_text_and_points_pictures_at_the_route():
    from research_lab.web.context import templates
    d = every_kind()
    html = S.render_html(d, lambda n: f"/x/asset/{n}", templates().env)
    assert "<script>alert(1)</script>" not in html and "&lt;script&gt;" in html
    assert html.count('class="dk-slide') == len(d.slides)
    assert 'src="/x/asset/cover.png"' in html and 'src="/x/asset/ours-0.png"' in html
    assert "dk-chip no" in html and "缺 glmer.nb" in html


def test_a_picture_the_deck_does_not_have_is_left_out_not_broken():
    from research_lab.web.context import templates
    d = every_kind()
    d.assets.pop("orig-1.png")
    html = S.render_html(d, lambda n: f"/a/{n}", templates().env)
    assert "/a/orig-1.png" not in html
    S.write_pptx(d, io.BytesIO())


def test_asset_names_are_plain_and_never_through_a_link(tmp_path):
    (tmp_path / "cover.png").write_bytes(b"x")
    os.symlink(tmp_path / "cover.png", tmp_path / "link.png")
    assert S.asset_file(tmp_path, "cover.png") == tmp_path / "cover.png"
    for bad in ("link.png", "../cover.png", "slides.json", "cover.PNG", "a/b.png", "", "x" * 90 + ".png"):
        assert S.asset_file(tmp_path, bad) is None


def test_a_chapter_is_read_into_slide_parts():
    did = method_notes.chapter("did")
    checks = S.chapter_checks(did.sections["稳健性检验：做什么，为什么"])
    assert checks[0][0] == "事件研究图" and checks[0][1] == "平行趋势、没有预期效应" and len(checks) == 6
    assert all(len(r[2]) <= 80 for r in checks)
    assumptions = S.chapter_assumptions(did.sections["识别假设"])
    assert [a for a, _ in assumptions][:3] == ["平行趋势", "没有预期效应", "不外溢"]
    assert S.chapter_formula(did.sections["模型设定"]).startswith("Y_it = α + β·Treat_i")
    reading = S.first_bullets(method_notes.chapter("poisson").sections["怎么读结果"], 2)
    assert reading[0].startswith("系数：") and "*" not in "".join(reading)


@pytest.mark.parametrize("key", sorted(method_notes.chapters()))
def test_every_chapter_gives_a_formula_and_a_checks_table(key):
    ch = method_notes.chapter(key)
    assert S.chapter_formula(ch.sections["模型设定"])
    assert len(S.chapter_checks(ch.sections["稳健性检验：做什么，为什么"])) >= 3
    assert S.first_bullets(ch.sections["怎么读结果"]) or S.first_paragraph(ch.sections["怎么读结果"])
